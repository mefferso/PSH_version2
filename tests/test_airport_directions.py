"""Tests for exact-peak direction recovery without direction guessing."""
import datetime as dt
import sys
import unittest
from pathlib import Path
from unittest.mock import patch
from collections import Counter
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import airport_fallback
from common import Audit,UTC

class DirectionTests(unittest.TestCase):
    def check(self,source_speed=15.432,source_time='2026-10-09T12:00:00Z',direction=80):
        w=Workbook();w.active.title='Wind and Pressure';s=w.active
        s.cell(2,11).value=30.0;s.cell(2,17).value=34;s.cell(2,23).value=1000
        audit=Audit();audit.add('Wind and Pressure',2,'KREG','wind',30,'kn',dt.datetime(2026,10,9,12,tzinfo=UTC),'https://example.org')
        payload={'SUMMARY':{'RESPONSE_CODE':1},'UNITS':{'wind_speed':'m/s','wind_direction':'Degrees'},
          'STATION':[{'STID':'KREG','SENSOR_VARIABLES':{'wind_speed':{'wind_speed_set_1':{'averaging_period_minutes':2}}},
          'OBSERVATIONS':{'date_time':[source_time],'wind_speed_set_1':[source_speed],'wind_direction_set_1':[direction]}}]}
        class Response:
            url='https://api.synopticdata.com/v2/stations/timeseries'
            def raise_for_status(self): pass
            def json(self):return payload
        qc=[];cnt=Counter()
        with patch.dict('os.environ',{'SYNOPTIC_TOKEN':'TEST'}),patch.object(airport_fallback,'inventory',return_value=[{'id':'KREG','row':2,'network':'AWOS'}]),patch.object(airport_fallback.requests,'get',return_value=Response()):
            airport_fallback.populate(w,qc,dt.date(2026,10,8),dt.date(2026,10,9),cnt,audit)
        return s,audit,qc,cnt
    def test_matching_peak_direction(self):
        s,a,q,c=self.check()
        self.assertEqual(s.cell(2,12).value,80)
        self.assertEqual(c['synoptic_airport_peak_directions_recovered'],1)
        self.assertEqual(a.entries[0]['direction_source'],'Synoptic exact-timestamp matched wind direction')
    def test_different_time_cannot_supply_direction(self):
        s,a,q,c=self.check(source_time='2026-10-09T12:01:00Z')
        self.assertIsNone(s.cell(2,12).value)
        self.assertEqual(c['synoptic_airport_peak_directions_recovered'],0)
    def test_different_speed_cannot_supply_direction(self):
        s,a,q,c=self.check(source_speed=10)
        self.assertIsNone(s.cell(2,12).value)
    def test_missing_direction_cannot_supply_direction(self):
        s,a,q,c=self.check(direction=None)
        self.assertIsNone(s.cell(2,12).value)
