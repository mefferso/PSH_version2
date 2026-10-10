import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from collections import Counter
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import coastal_water as cw
from common import UTC,Audit,inventory
ROOT=Path(__file__).resolve().parents[1]
A=dt.datetime(2026,10,8,tzinfo=UTC);B=A+dt.timedelta(days=2)

class CoastalWaterTests(unittest.TestCase):
    def test_all_31_exact_template_mappings_including_missing_wcc_link(self):
        wb=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
        stations=[s for s in inventory(wb,'Water Level') if s['network'] in ('USACE','LA CPRA')]
        mapping=cw.load_mapping()
        self.assertEqual(set(mapping),{s['id'] for s in stations});self.assertEqual(len(mapping),31)
        self.assertEqual(mapping['WCCL1']['rivergages_sid'],'76265')
        self.assertEqual(mapping['WCCL1']['cwms_location'],'WestCC _FS')
        self.assertNotEqual(mapping['CMPL1']['cwms_location'],mapping['85750LA']['cwms_location'])
        for s in stations:self.assertTrue(cw.match_template(s,wb,mapping[s['id']]))
        s=stations[0];wb['Water Level'].cell(s['row'],3).value=0
        self.assertFalse(cw.match_template(s,wb,mapping[s['id']]))

    def payload(self,values,units='ft'):
        return {'name':'X.Stage.Inst.1Hour.0.rev','office-id':'MVN','units':units,
                'value-columns':[{'name':'date-time','ordinal':1},{'name':'value','ordinal':2},{'name':'quality-code','ordinal':3}],
                'values':values,'total':len(values)}

    def test_parser_preserves_original_and_partial_peak_and_utc_half_open_window(self):
        ms=int(A.timestamp()*1000)
        p=self.payload([[ms,1,3],[ms+3600000,None,9],[ms+7200000,2,0],[int(B.timestamp()*1000),999,3]],'m')
        rows=cw.parse_cwms(p,p['name'],A,B)
        self.assertEqual(len(rows),3)
        peak=cw.summarize(rows,A,B,3600)['peak']
        self.assertEqual(peak['original_value'],2);self.assertEqual(peak['original_unit'],'m')
        self.assertAlmostEqual(peak['value_ft'],2/0.3048)
        self.assertEqual(peak['time_utc'],'2026-10-08T02:00:00+00:00')
        self.assertFalse(cw.summarize(rows,A,B,3600)['complete'])

    def test_quality_missing_rejected_invalid_and_conflicting_duplicate_do_not_create_false_peak(self):
        ms=int(A.timestamp()*1000)
        p=self.payload([[ms,1,3],[ms+3600000,999,17],[ms+7200000,-9999,9],[ms+10800000,'NaN',0],
                        [ms+14400000,5,3],[ms+14400000,6,3]])
        result=cw.summarize(cw.parse_cwms(p,p['name'],A,B),A,B,3600)
        self.assertEqual(result['peak']['value_ft'],1)
        self.assertFalse(result['complete']);self.assertEqual(result['conflicting_times'],1)

    def test_units_identity_and_schema_are_explicit(self):
        for change in [{'name':'wrong'},{'office-id':'NWO'},{'units':'unknown'},{'value-columns':[]}]:
            p=self.payload([]);p.update(change)
            with self.subTest(change=change),self.assertRaises(ValueError):cw.parse_cwms(p,'X.Stage.Inst.1Hour.0.rev',A,B)

    def test_pagination_keeps_previous_observations_when_later_page_fails(self):
        ms=int(A.timestamp()*1000);p=self.payload([[ms,2,3]]);p.update({'total':2,'next-page':'opaque'})
        response=Mock();response.json.return_value=p;response.url='https://cwms-data.usace.army.mil/cwms-data/timeseries?name=X'
        session=Mock();session.get.side_effect=[response,cw.requests.Timeout()]
        result=cw.collect({'cwms_timeseries':p['name'],'interval_seconds':3600},A,B,session=session)
        self.assertEqual(result['coverage']['peak']['value_ft'],2)
        self.assertFalse(result['coverage']['complete']);self.assertTrue(result['errors'])
        self.assertEqual(len(result['source_urls']),1)

    def test_location_altitude_and_requested_datum_are_not_measurement_evidence(self):
        mapping=cw.load_mapping()
        meta={'vertical-datum':'NAVD88','elevation':0,'description':'Gage location'}
        self.assertFalse(cw.qualify(mapping['TSPL1'],meta,A,B)[0])
        self.assertFalse(cw.qualify(mapping['SWBL1'],meta,A,B)[0])
        self.assertFalse(cw.qualify(mapping['BBOL1'],meta,A,B)[0])
        self.assertFalse(cw.qualify(mapping['COCL1'],meta,A,B)[0])

    def test_direct_station_datum_requires_effective_period_and_live_description(self):
        m=cw.load_mapping()['76305']
        meta={'description':m['description'],'vertical-datum':m['location_vertical_datum']}
        self.assertTrue(cw.qualify(m,meta,A,B)[0])
        self.assertFalse(cw.qualify(m,dict(meta,description='Gage reset to new datum'),A,B)[0])
        self.assertFalse(cw.qualify(m,meta,dt.datetime(2010,1,1,tzinfo=UTC),A)[0])

    def test_partial_qualified_and_unqualified_peaks_survive_in_review_outputs(self):
        wb=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx');q=wb.create_sheet('QC');audit=Audit()
        mapping=cw.load_mapping();ms=int(A.timestamp()*1000)
        def fetch(m,a,b,**kwargs):
            p=self.payload([[ms,2,3]]);p['name']=m['cwms_timeseries']
            rows=cw.parse_cwms(p,p['name'],a,b)
            return {'observations':rows,'coverage':cw.summarize(rows,a,b,m['interval_seconds']),
                    'source_urls':['https://cwms-data.usace.army.mil/cwms-data/timeseries?name='+m['cwms_timeseries']],
                    'errors':[],'pages':[]}
        locations={m['cwms_location']:{'name':m['cwms_location'],'office-id':'MVN','public-name':m['public_name'],
                     'latitude':m['source_latitude'],'longitude':m['source_longitude'],
                     'description':m['description'],'vertical-datum':m['location_vertical_datum']} for m in mapping.values()}
        with tempfile.TemporaryDirectory() as d,patch.object(cw,'collect',side_effect=fetch),patch.object(cw,'fetch_locations',return_value=locations),patch.dict('os.environ',{},clear=True):
            cw.populate(wb,q,A.date(),(B-dt.timedelta(days=1)).date(),Counter(),audit,output_dir=Path(d))
            reports=json.loads((Path(d)/'coastal_water_audit.json').read_text());by_id={x['site_id']:x for x in reports}
            self.assertEqual(len(reports),31)
            self.assertTrue(by_id['76305']['can_populate_psh']);self.assertFalse(by_id['TSPL1']['can_populate_psh'])
            self.assertEqual(by_id['TSPL1']['peak_ft'],2);self.assertEqual(by_id['TSPL1']['flag'],'I')
            self.assertIn('Water Level Review',wb.sheetnames)
            s=next(s for s in inventory(wb,'Water Level') if s['id']=='76305')
            self.assertEqual(wb['Water Level'].cell(s['row'],7).value,2);self.assertEqual(wb['Water Level'].cell(s['row'],14).value,'I')
            t=next(s for s in inventory(wb,'Water Level') if s['id']=='TSPL1')
            self.assertIsNone(wb['Water Level'].cell(t['row'],7).value)
            self.assertTrue(all(e['datum']=='NAVD88' and e['datum_evidence'] for e in audit.entries))

if __name__=='__main__':unittest.main()

class CoastalIntegrationTests(unittest.TestCase):
    def test_real_build_validates_review_outputs_and_water_tab_shows_unqualified_peak(self):
        import os,build,validate,export_dashboard
        def adapter(w,q,a,b,c,audit):
            reports=[{'site_id':'TSPL1','row':24,'network':'USACE','rivergages_sid':'85300',
                      'cwms_timeseries':'Tickfaw_River-Springfield.Stage.Inst.1Hour.0.rev','peak_ft':4.22,
                      'peak_time_utc':'2026-10-09T22:00:00+00:00','observed_datum':'NGVD29 stage',
                      'can_populate_psh':False,'human_review':True,'flag':'I','observation_count':1,
                      'coverage_complete':False,'reason':'Controlled incomplete stage fixture',
                      'observation_url':'https://cwms-data.usace.army.mil/cwms-data/locations/Tickfaw_River-Springfield?office=MVN',
                      'source_url':'https://cwms-data.usace.army.mil/cwms-data/timeseries?name=Tickfaw'}]
            cw.write_outputs(w,build.OUT,reports,{})
        with tempfile.TemporaryDirectory() as d,patch.dict(os.environ,{'STORM_NAME':'Hurricane Isaias','START_UTC':'2026-10-08','END_UTC':'2026-10-09','PSH_OFFLINE':'1'},clear=True):
            out=Path(d)/'output';site=Path(d)/'site'
            with patch.object(build,'OUT',out),patch.object(build,'TEMPLATE',ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'),patch.object(build,'collectors',return_value=[adapter]):
                meta=build.build()
            self.assertTrue(validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'))
            payload=export_dashboard.export(out,site)
            rows=payload['tabs']['Water Level']['rows'];tick=next(r for r in rows if r['v'][0]=='TSPL1')
            self.assertIsNone(tick['v'][6]);self.assertIn(4.22,tick['v'][14:])
            usgs_tick=next(r for r in rows if r['v'][0]=='TSPL1' and r['row']==71)
            self.assertTrue(all(v is None for v in usgs_tick['v'][14:]))
            self.assertIn('Review peak (ft; original datum)',payload['tabs']['Water Level']['headers'])
            validate.validate_site(out,site)
            self.assertTrue((site/'coastal_water_review.csv').exists())

    def test_branch_dispatch_and_pr_cannot_publish(self):
        workflow=(ROOT/'.github/workflows/build-psh.yml').read_text()
        self.assertIn("github.ref == 'refs/heads/main'",workflow)
        self.assertIn("github.event_name != 'pull_request'",workflow)

class TemplateLinkTests(unittest.TestCase):
    def test_link_patch_changes_only_water_hyperlinks_and_adds_missing_wcc(self):
        import zipfile
        import repair_coastal_links
        source=ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'
        with tempfile.TemporaryDirectory() as d:
            target=Path(d)/'template.xlsx'
            repair_coastal_links.repair(source,target)
            with zipfile.ZipFile(source) as old,zipfile.ZipFile(target) as new:
                self.assertEqual(old.namelist(),new.namelist())
                changed=[n for n in old.namelist() if old.read(n)!=new.read(n)]
                self.assertTrue(set(changed).issubset({'xl/worksheets/sheet4.xml','xl/worksheets/_rels/sheet4.xml.rels'}))
            w=load_workbook(target);mapping=cw.load_mapping()
            for s in inventory(w,'Water Level'):
                if s['id'] in mapping and s['network'] in ('USACE','LA CPRA'):self.assertEqual(s['url'],mapping[s['id']]['observation_url'])
            old=load_workbook(source)
            for tab in old.sheetnames:
                self.assertEqual(old[tab].max_row,w[tab].max_row)
                for arow,brow in zip(old[tab],w[tab]):
                    for a,b in zip(arow,brow):
                        self.assertEqual(a.value,b.value)
                        self.assertEqual(a._style,b._style)

class FallbackTests(unittest.TestCase):
    def test_missing_cwms_peak_uses_only_original_exact_id_hml_and_preserves_secondary_provenance(self):
        m=cw.load_mapping()['MBBL1']
        no={'observations':[],'coverage':cw.summarize([],A,B,3600),'source_urls':[],'errors':[],'pages':[]}
        response=Mock();response.text='station,valid[UTC],Stage[ft]\nMBBL1,2026-10-08 12:00,1.95\nOTHER,2026-10-08 12:00,99\n';response.url='https://mesonet.agron.iastate.edu/cgi-bin/request/hml.py?station=MBBL1'
        s=Mock();s.get.return_value=response
        result=cw.collect_hml(m,A,B,session=s)
        self.assertEqual(result['coverage']['peak']['original_value'],1.95)
        self.assertEqual(result['coverage']['peak']['time_utc'],'2026-10-08T12:00:00+00:00')
        self.assertEqual(result['source_kind'],'NWS HML archived by IEM (secondary)')
        self.assertFalse(result['coverage']['complete'])
        with patch.object(cw,'collect',return_value=no),patch.object(cw,'collect_hml',return_value=result):
            recovered=cw.retrieve(m,A,B)
        self.assertEqual(recovered['coverage']['peak']['value_ft'],1.95)
        self.assertFalse(recovered['datum_eligible'])

class CapturedCWMSContractTests(unittest.TestCase):
    def test_real_partial_surge_barrier_observations_keep_measured_peak(self):
        p=json.loads((ROOT/'tests/fixtures/cwms_SBEL1_20261008-10.json').read_text())
        result=cw.summarize(cw.parse_cwms(p,p['name'],A,B),A,B,900)
        self.assertEqual(result['count'],179);self.assertFalse(result['complete'])
        self.assertAlmostEqual(result['peak']['value_ft'],4.19)
        self.assertEqual(result['peak']['time_utc'],'2026-10-09T20:30:00+00:00')
    def test_real_cocodrie_series_uses_ft_and_actual_utc_observation_times(self):
        p=json.loads((ROOT/'tests/fixtures/cwms_76305_20261008-10.json').read_text())
        result=cw.summarize(cw.parse_cwms(p,p['name'],A,B),A,B,3600)
        self.assertTrue(result['complete']);self.assertEqual(result['count'],48)
        self.assertAlmostEqual(result['peak']['value_ft'],2.55)
        self.assertEqual(result['peak']['time_utc'],'2026-10-09T21:00:00+00:00')

class PagesVerificationTests(unittest.TestCase):
    def test_deployed_manifest_and_water_tab_must_match_validated_export(self):
        import verify_pages
        local={'meta':{'commit':'new'},'tabs':{'Water Level':{'headers':['ID','Review peak'],'rows':[{'v':['SBEL1',4.19]}]}},'water_review':[{'site_id':'SBEL1','peak_ft':4.19}]}
        self.assertTrue(verify_pages.matches(local,json.loads(json.dumps(local))))
        old=json.loads(json.dumps(local));old['meta']['commit']='old'
        self.assertFalse(verify_pages.matches(local,old))
        wrong=json.loads(json.dumps(local));wrong['tabs']['Water Level']['rows'][0]['v'][1]=None
        self.assertFalse(verify_pages.matches(local,wrong))

class LegacyLinkCompatibilityTests(unittest.TestCase):
    def test_existing_rivergages_adapter_recognizes_only_audited_cwms_links(self):
        import usace
        cell=Mock();m=cw.load_mapping()['BBOL1'];cell.hyperlink.target=m['observation_url']
        self.assertEqual(usace.station_link(cell)[0],'52800')
        cell.hyperlink.target=m['observation_url'].replace('Bayou_Boeuf-Amelia','OTHER')
        self.assertIsNone(usace.station_link(cell))
