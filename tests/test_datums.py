import datetime as dt
import sys
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import datums
from common import UTC
class DatumTests(unittest.TestCase):
    def test_offset_requires_identity_independent_evidence_and_effective_window(self):
        a=dt.datetime(2024,9,11,tzinfo=UTC);b=a+dt.timedelta(days=1)
        st={'id':'BSGL1','row':10,'url':'https://rivergages.mvr.usace.army.mil/WaterControl/stationinfo2.cfm?sid=82742'}
        x={'site_id':'BSGL1','rivergages_sid':'82742','datum':'NAVD88','offset_ft':0.11,'reviewed':True,
            'effective_start_utc':'2024-01-01T00:00:00Z','effective_end_utc':'2025-01-01T00:00:00Z',
            'evidence':['https://rivergages.mvr.usace.army.mil/WaterControl/stationinfo2.cfm?sid=82742','https://www.weather.gov/lix/']}
        self.assertAlmostEqual(datums.offset(x,st,a,b),0.11)
        for change in [{'rivergages_sid':'76030'},{'reviewed':False},{'effective_end_utc':'2024-09-11T12:00:00Z'},{'evidence':x['evidence'][:1]}]:
            with self.subTest(change=change),self.assertRaises(ValueError):datums.offset(dict(x,**change),st,a,b)
    def test_hml_requires_explicit_utc_and_feet(self):
        a=dt.date(2024,9,11)
        rows=datums.parse_hml('station,valid[utc],stage[ft]\nBSGL1,2024-09-11 12:00,2.5\nOTHER,2024-09-11 12:00,30\nBSGL1,2024-09-12 12:00,20\n','BSGL1',a,a)
        self.assertEqual([x[0] for x in rows],[2.5])
        with self.assertRaises(ValueError):datums.parse_hml('station,valid,stage\nBSGL1,2024-09-11 12:00,2.5\n','BSGL1',a,a)

class LiveHMLContractTests(unittest.TestCase):
    def test_captured_archive_capitalized_headers_have_explicit_utc_and_feet(self):
        text=(Path(__file__).resolve().parent/'fixtures/iem_hml_BBOL1_20240910-12.csv').read_text()
        rows=datums.parse_hml(text,'BBOL1',dt.date(2024,9,10),dt.date(2024,9,12))
        self.assertEqual(len(rows),72);self.assertEqual(max(rows)[0],4.4)
    def test_raw_historical_stage_is_retrieved_but_not_written_without_a_datum_registry(self):
        from openpyxl import load_workbook
        from collections import Counter
        from unittest.mock import patch
        import os
        from common import Audit,inventory
        w=load_workbook(Path(__file__).resolve().parents[1]/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx');q=w.create_sheet('QC');audit=Audit()
        station=next(x for x in inventory(w,'Water Level') if x['id']=='BBOL1')
        def fetch(sid,start,end):return ([(4.4,dt.datetime(2024,9,12,2,tzinfo=UTC))] if sid=='BBOL1' else []),'https://mesonet.agron.iastate.edu/cgi-bin/request/hml.py?station='+sid
        with patch.dict(os.environ,{},clear=True),patch.object(datums,'collect',side_effect=fetch):datums.populate(w,q,dt.date(2024,9,10),dt.date(2024,9,12),Counter(),audit)
        self.assertIsNone(w['Water Level'].cell(station['row'],7).value)
        self.assertEqual(audit.entries,[])
        self.assertTrue(any(row[1]=='BBOL1' and '4.4' in str(row[4]) for row in q.iter_rows(values_only=True)))
