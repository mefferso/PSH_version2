import csv
import datetime as dt
import math
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import common
import products

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'

class ProductTests(unittest.TestCase):
    def test_inventory_excludes_footer_preserves_numeric_ids(self):
        w = load_workbook(TEMPLATE)
        rows = list(common.inventory(w, 'Wind and Pressure'))
        self.assertEqual(len(rows), 169)
        self.assertIn('42040', [x['id'] for x in rows])
        self.assertNotIn('Latest Update:', [x['id'] for x in rows])

    def test_summary_sorts_and_excludes_high_or_unknown_anemometers(self):
        w = load_workbook(TEMPLATE)
        s = w['Wind and Pressure']
        for row, value, height in [(2, 60, 10), (3, 70, 10), (4, 100, 21), (5, 120, None)]:
            s.cell(row, 11, value); s.cell(row, 9).value = height
        products.summaries(w)
        self.assertEqual(w['Summary']['D20'].value, 70)
        self.assertEqual(w['Summary']['D21'].value, 60)
        self.assertIsNone(w['Summary']['D22'].value)
        self.assertFalse(any(c.data_type == 'f' for row in w['Summary'] for c in row))

    def test_csv_excludes_duplicate_ids_and_empty_measurements(self):
        w = load_workbook(TEMPLATE)
        for r in range(2,171):
            for c in (11,17,23): w['Wind and Pressure'].cell(r,c).value=None
        for r in [145,155]: w['Wind and Pressure'].cell(r,11,45)
        with tempfile.TemporaryDirectory() as d:
            issues = products.export_csv(w, Path(d))
            with open(Path(d)/'Wind_and_Pressure_REVIEW.csv',newline='') as f: rows=list(csv.reader(f))
            self.assertEqual(len(rows),1)
            self.assertTrue(any('BYGL1' in i for i in issues))

    def test_audit_rejects_nonfinite_and_unqualified_water(self):
        audit=common.Audit()
        with self.assertRaises(ValueError):
            audit.add('Water Level',2,'X','water',math.nan,'ft',dt.datetime.now(dt.timezone.utc),'https://example.com',datum='NAVD88')
        with self.assertRaises(ValueError):
            audit.add('Water Level',2,'X','water',2,'ft',dt.datetime.now(dt.timezone.utc),'https://example.com')

    def test_safe_total_requires_complete_nonoverlapping_intervals(self):
        z=dt.timezone.utc; a=dt.datetime(2024,9,11,tzinfo=z);b=a+dt.timedelta(hours=2)
        readings=[(a,a+dt.timedelta(hours=1),0.2),(a+dt.timedelta(hours=1),b,0.4)]
        self.assertAlmostEqual(common.interval_total(readings,a,b),0.6)
        self.assertIsNone(common.interval_total(readings[:1],a,b))
        self.assertIsNone(common.interval_total(readings+[readings[0]],a,b))
        self.assertIsNone(common.interval_total([(a,b,None)],a,b))

if __name__=='__main__':unittest.main()

class ReportingThresholdTests(unittest.TestCase):
    def test_current_nws_thresholds_without_dropping_review_readings(self):
        w=load_workbook(TEMPLATE)
        for st in common.inventory(w,'Wind and Pressure'):
            for c in (11,17,23):w['Wind and Pressure'].cell(st['row'],c).value=None
        w['Wind and Pressure'].cell(2,17).value=33
        w['Wind and Pressure'].cell(3,17).value=33.1
        w['Wind and Pressure'].cell(4,23).value=1004.9
        w['Wind and Pressure'].cell(5,23).value=1005
        w['Rainfall'].cell(2,8).value=3
        w['Rainfall'].cell(3,8).value=2.99
        with tempfile.TemporaryDirectory() as d:
            products.export_csv(w,Path(d))
            with open(Path(d)/'WindandPressure_CANDIDATE.csv',newline='') as f:rows=list(csv.reader(f))
            self.assertEqual({x[0] for x in rows[1:]},{'KMSY','KNEW'})
            self.assertEqual(len(rows[0]),29)
            with open(Path(d)/'Rainfall_CANDIDATE.csv',newline='') as f:rows=list(csv.reader(f))
            self.assertEqual({x[0] for x in rows[1:]},{'BTR'})
