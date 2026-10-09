import datetime as dt
import sys
import unittest
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import imports
from common import Audit
ROOT=Path(__file__).resolve().parents[1]
class ImportTests(unittest.TestCase):
    def setUp(self):
        self.w=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
        self.audit=Audit();self.start=dt.date(2024,9,11);self.end=self.start
    def reading(self,**kwargs):
        result={'tab':'Wind and Pressure','row':91,'site_id':'XBTR','network':'WeatherFlow',
            'variable':'gust','value':52.5,'unit':'kn','time_utc':'2024-09-11T12:00:00Z',
            'source_url':'https://ds.weatherflow.com/map','reviewed':True}
        result.update(kwargs);return result
    def test_manual_weatherflow_exact_row_and_time(self):
        imports.apply(self.w,[self.reading()],self.start,self.end,self.audit)
        self.assertEqual(self.w['Wind and Pressure'].cell(91,17).value,52.5)
        self.assertEqual(self.w['Wind and Pressure'].cell(91,19).value,'1200')
        self.assertEqual(len(self.audit.entries),1)
    def test_identity_window_units_and_unreviewed_reject(self):
        for change in [{'site_id':'MSY'},{'time_utc':'2024-09-12T12:00:00Z'}, {'unit':'mph'}, {'reviewed':False}]:
            with self.subTest(change=change),self.assertRaises(ValueError):
                imports.apply(self.w,[self.reading(**change)],self.start,self.end,self.audit)
    def test_stage_cannot_be_relabelled_without_two_datum_evidence_sources(self):
        reading=self.reading(tab='Water Level',row=10,site_id='BSGL1',network='LA CPRA',
            variable='water',value=3.2,unit='ft',datum='NAVD88',
            source_url='https://rivergages.mvr.usace.army.mil/WaterControl/stationinfo2.cfm?sid=82742')
        with self.assertRaises(ValueError):imports.apply(self.w,[reading],self.start,self.end,self.audit)
        reading['datum_evidence']=['https://rivergages.mvr.usace.army.mil/WaterControl/stationinfo2.cfm?sid=82742','https://www.weather.gov/lix/']
        imports.apply(self.w,[reading],self.start,self.end,self.audit)
        self.assertEqual(self.w['Water Level'].cell(10,8).value,'NAVD88')
    def test_rainfall_must_cover_explicit_window(self):
        reading=self.reading(tab='Rainfall',row=98,site_id='LA-AS-02',network='CoCoRaHS',
            variable='rain',value=4.2,unit='in',time_utc='2024-09-12T00:00:00Z',
            source_url='https://www.cocorahs.org/ViewData/ListDailyPrecipReports.aspx')
        with self.assertRaises(ValueError):imports.apply(self.w,[reading],self.start,self.end,self.audit)
        reading['interval_start_utc']='2024-09-11T00:00:00Z'
        imports.apply(self.w,[reading],self.start,self.end,self.audit)
        self.assertEqual(self.w['Rainfall'].cell(98,8).value,4.2)
