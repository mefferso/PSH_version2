import hashlib
import sys
import unittest
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import regression
ROOT=Path(__file__).resolve().parents[1]
REFERENCE=ROOT/'tests/fixtures/PSHLIX_2024AL06_Francine_Data.xlsx'
class ReferenceTests(unittest.TestCase):
    def test_issued_reference_bytes_are_unchanged(self):
        self.assertEqual(hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),'7347750b44507e7e9025baa25c9db68a7aa686868d6d7d1baee0e03fa468c16f')
    def test_no_generated_readings_do_not_count_as_accuracy(self):
        generated=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
        import build
        build.reset(generated)
        result=regression.compare(generated,load_workbook(REFERENCE))
        self.assertGreater(result['missing'],0)
        self.assertEqual(result['compared'],0)
        self.assertIsNone(result['agreement'])
    def test_actual_issued_reference_detects_numerical_discrepancy(self):
        generated=load_workbook(REFERENCE);s=generated['Wind and Pressure']
        target=next(r for r in range(2,s.max_row+1) if s.cell(r,1).value=='KBTR')
        expected=s.cell(target,17).value;s.cell(target,17).value=expected+5
        result=regression.compare(generated,load_workbook(REFERENCE))
        self.assertTrue(any(x['site_id']=='KBTR' and x['variable']=='gust' and x['status']=='DISCREPANCY' for x in result['observations']))

class RainWindowRegressionTests(unittest.TestCase):
    def test_different_rainfall_intervals_are_not_numerical_accuracy_passes(self):
        from copy import deepcopy
        import regression
        reference=load_workbook(ROOT/'tests/fixtures/PSHLIX_2024AL06_Francine_Data.xlsx',data_only=True)
        generated=deepcopy(reference);generated['Rainfall']['B228']='0000 UTC Sep 10 2024'
        result=regression.compare(generated,reference)
        rain=[x for x in result['observations'] if x['variable']=='rain' and x['generated'] is not None]
        self.assertTrue(rain);self.assertTrue(all(x['status']=='WINDOW MISMATCH' for x in rain))
        self.assertEqual(result['window_mismatch'],len(rain))
