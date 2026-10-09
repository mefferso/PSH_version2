import datetime as dt
import sys
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import tornadoes
class TornadoTests(unittest.TestCase):
    def test_confirmed_storm_association_cwa_and_standard_time_to_utc(self):
        row={'EVENT_TYPE':'Tornado','WFO':'LIX','EPISODE_NARRATIVE':'Hurricane Francine affected the area.',
             'EVENT_NARRATIVE':'Survey confirmed tornado.','BEGIN_DATE_TIME':'11-SEP-24 18:30:00','CZ_TIMEZONE':'CST-6',
             'BEGIN_LAT':'30.2','BEGIN_LON':'-90.2','BEGIN_LOCATION':'NEW ORLEANS','CZ_NAME':'ORLEANS',
             'STATE':'LOUISIANA','TOR_F_SCALE':'EF1','EVENT_ID':'123'}
        events=tornadoes.parse([row],'Hurricane Francine',dt.date(2024,9,11),dt.date(2024,9,12))
        self.assertEqual(len(events),1)
        self.assertEqual(events[0]['time'].isoformat(),'2024-09-12T00:30:00+00:00')
        row['EPISODE_NARRATIVE']='Thunderstorms during a cold front.'
        self.assertEqual(tornadoes.parse([row],'Hurricane Francine',dt.date(2024,9,11),dt.date(2024,9,12)),[])
        row['EPISODE_NARRATIVE']='Francine';row['CZ_TIMEZONE']='unknown'
        self.assertEqual(tornadoes.parse([row],'Hurricane Francine',dt.date(2024,9,11),dt.date(2024,9,12)),[])

class TornadoPublicationTests(unittest.TestCase):
    def test_no_named_records_does_not_assert_zero_and_overflow_is_atomic(self):
        from openpyxl import Workbook
        from collections import Counter
        from unittest.mock import patch
        import os
        from common import Audit
        w=Workbook();w.active.title='Summary';s=w.create_sheet('Tornadoes');s.append(['header']*11);q=w.create_sheet('QC');audit=Audit()
        with patch.dict(os.environ,{'STORM_NAME':'Hurricane Francine'}),patch.object(tornadoes,'collect',return_value=([],[tornadoes.BASE])):
            tornadoes.populate(w,q,dt.date(2024,9,11),dt.date(2024,9,12),Counter(),audit)
        self.assertIsNone(w['Summary']['B11'].value)
        event={'id':'1','time':dt.datetime(2024,9,11,12,tzinfo=dt.timezone.utc),'lat':30,'lon':-90,'rating':'EF1','city':'X','county':'Y','state':'LA','narrative':'Francine','zone':'CST-6','url':tornadoes.BASE}
        with patch.dict(os.environ,{'STORM_NAME':'Hurricane Francine'}),patch.object(tornadoes,'collect',return_value=([event],[tornadoes.BASE])):
            tornadoes.populate(w,q,dt.date(2024,9,11),dt.date(2024,9,12),Counter(),audit)
        self.assertEqual(audit.entries,[]);self.assertEqual(s.max_row,1)
