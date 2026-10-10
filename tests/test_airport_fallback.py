"""Airport Synoptic archive fallback must never override primary measurements."""
import datetime as dt
import sys
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import airport_fallback
from common import Audit

class Response:
    url='https://api.synopticdata.com/v2/stations/timeseries?stid=KREG&token=PRIVATE'
    def raise_for_status(self): pass
    def json(self): return self.payload

class AirportFallbackTests(unittest.TestCase):
    def setUp(self):
        self.wb=Workbook();self.wb.active.title='Wind and Pressure'
        self.stations=[{'id':'KREG','row':2,'network':'AWOS','url':'https://example.org'}]
        self.start=dt.date(2026,10,8);self.end=dt.date(2026,10,9)
    def run_fallback(self,payload):
        reply=Response();reply.payload=payload
        qc=[];audit=Audit();counts=Counter()
        with patch.dict('os.environ',{'SYNOPTIC_TOKEN':'PRIVATE'}),patch.object(airport_fallback,'inventory',return_value=self.stations),patch.object(airport_fallback.requests,'get',return_value=reply) as request:
            airport_fallback.populate(self.wb,qc,self.start,self.end,counts,audit)
        return qc,audit,counts,request
    def payload(self,stid='KREG',period=2):
        sensor={'wind_speed':{'wind_speed_set_1':{'averaging_period_minutes':period}}} if period is not None else {}
        return {'SUMMARY':{'RESPONSE_CODE':1},
                'UNITS':{'wind_speed':'m/s','wind_gust':'m/s','wind_direction':'Degrees','sea_level_pressure':'hPa'},
                'STATION':[{'STID':stid,'SENSOR_VARIABLES':sensor,'OBSERVATIONS':{
                'date_time':['2026-10-09T12:00:00Z'],
                'wind_speed_set_1':[12],'wind_gust_set_1':[19],
                'sea_level_pressure_set_1':[1000],'wind_direction_set_1':[80]}}]}
    def test_fills_missing_qualified_fields_and_preserves_existing_wind(self):
        self.wb['Wind and Pressure'].cell(2,11).value=31
        qc,audit,counts,request=self.run_fallback(self.payload())
        row=self.wb['Wind and Pressure']
        self.assertEqual(row.cell(2,11).value,31)
        self.assertAlmostEqual(row.cell(2,17).value,round(19*1.9438444924406,1))
        self.assertEqual(row.cell(2,23).value,1000)
        self.assertEqual(len(audit.entries),2)
        self.assertEqual(counts['synoptic_airport_fields_recovered'],2)
        self.assertEqual(request.call_args.kwargs['params']['stid'],'KREG')
        self.assertEqual(request.call_args.kwargs['params']['sensorvars'],1)
        self.assertNotIn('PRIVATE',audit.entries[0]['source_url'])
    def test_exact_station_not_substituted(self):
        qc,audit,counts,_=self.run_fallback(self.payload('KHDC'))
        self.assertFalse(audit.entries)
        self.assertEqual(qc[0][3],'NO DATA')
    def test_unknown_wind_averaging_withheld_but_gust_pressure_retained(self):
        qc,audit,counts,_=self.run_fallback(self.payload(period=None))
        self.assertIsNone(self.wb['Wind and Pressure'].cell(2,11).value)
        self.assertIsNotNone(self.wb['Wind and Pressure'].cell(2,17).value)
        self.assertIsNotNone(self.wb['Wind and Pressure'].cell(2,23).value)
    def test_station_altimeter_not_treated_as_mslp(self):
        payload=self.payload()
        payload['STATION'][0]['OBSERVATIONS'].pop('sea_level_pressure_set_1')
        payload['STATION'][0]['OBSERVATIONS']['altimeter_set_1']=[29.85]
        qc,audit,counts,_=self.run_fallback(payload)
        self.assertIsNone(self.wb['Wind and Pressure'].cell(2,23).value)
