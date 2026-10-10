"""Regression for IEM higher METAR gust replacing minute-archive gust."""
import datetime as dt
import sys
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import iem
from common import Audit,UTC

class GustArchivePrecedenceTests(unittest.TestCase):
    def test_metar_peak_gust_replaces_lower_minute_gust_and_audit(self):
        wb=Workbook();s=wb.active;s.title='Wind and Pressure'
        wb.create_sheet('Rainfall')
        s.cell(2,17).value=25
        t=dt.datetime(2026,10,9,15,tzinfo=UTC)
        audit=Audit()
        audit.add('Wind and Pressure',2,'KPQL','gust',25,'kn',t,'https://example.org/minute')
        good={'time':t+dt.timedelta(minutes=3),'wind':None,'gust':47,'gust_dir':180,
              'pressure':None,'source_kind':'METAR PK WND remark','raw':{'metar':'PK WND 18047/1503'}}
        stations=[{'id':'KPQL','network':'ASOS','row':2}]
        with patch.object(iem,'inventory',side_effect=lambda wb,tab:stations if tab=='Wind and Pressure' else []),patch.object(iem,'collect',return_value=([good],'https://example.org/iem')):
            iem.populate(wb,[],dt.date(2026,10,8),dt.date(2026,10,9),Counter(),audit)
        self.assertEqual(s.cell(2,17).value,47)
        self.assertEqual(s.cell(2,18).value,180)
        found=[x for x in audit.entries if x['variable']=='gust']
        self.assertEqual(len(found),1)
        self.assertEqual(found[0]['value'],47)
        self.assertEqual(found[0]['source_url'],'https://example.org/iem')
    def test_lower_metar_gust_does_not_replace_higher_minute_gust(self):
        wb=Workbook();s=wb.active;s.title='Wind and Pressure';wb.create_sheet('Rainfall')
        s.cell(2,17).value=47
        t=dt.datetime(2026,10,9,15,tzinfo=UTC)
        audit=Audit();audit.add('Wind and Pressure',2,'KPQL','gust',47,'kn',t,'https://example.org/minute')
        rows=[{'time':t,'wind':None,'gust':25,'pressure':None}]
        stations=[{'id':'KPQL','network':'ASOS','row':2}]
        with patch.object(iem,'inventory',side_effect=lambda wb,tab:stations if tab=='Wind and Pressure' else []),patch.object(iem,'collect',return_value=(rows,'https://example.org/iem')):
            iem.populate(wb,[],dt.date(2026,10,8),dt.date(2026,10,9),Counter(),audit)
        self.assertEqual(s.cell(2,17).value,47)
        self.assertEqual([x['value'] for x in audit.entries if x['variable']=='gust'],[47])
