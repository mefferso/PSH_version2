"""Archived Francine sustained winds: source selection, not reference fitting."""
import csv,datetime as dt,io,json,sys,unittest,zipfile
from pathlib import Path
from collections import Counter
from unittest.mock import patch
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import iem,asos1min
from common import Audit,UTC
ROOT=Path(__file__).resolve().parents[1]
EVIDENCE=ROOT/'docs/validation/francine-asos-20261010/source-records.zip'
class FrancineASOSTests(unittest.TestCase):
 def test_independent_noaa_fixed_fields_match_iem(self):
  with zipfile.ZipFile(EVIDENCE) as z:
   for sid,expected,timestamp in [('KMSY',50,'2024-09-12T02:01:00'),('KNEW',42,'2024-09-12T02:02:00'),('KPQL',23,'2024-09-12T05:25:00')]:
    good=[]
    for line in z.read(sid+'-noaa.dat').decode().splitlines():
     try:
      local=dt.datetime.strptime(line[13:25],'%Y%m%d%H%M')
      utc=local.replace(hour=int(line[25:27]),minute=int(line[27:29]))
      if utc.hour<local.hour:utc+=dt.timedelta(days=1)
      if dt.datetime(2024,9,10)<=utc<dt.datetime(2024,9,13):good.append((int(line[74:79]),utc))
     except ValueError:continue
    peak=max(v for v,t in good)
    self.assertEqual(peak,expected)
    self.assertEqual([t.isoformat() for v,t in good if v==peak],[timestamp])
 def test_archived_sources_and_exact_peaks(self):
  with zipfile.ZipFile(EVIDENCE) as z:
   for sid,minute,aviation in [('KBTR',None,27),('KMSY',50,43),('KNEW',42,37),('KPQL',23,20)]:
    rows=asos1min.parse(z.read(sid+'-minute.csv').decode(),sid,dt.date(2024,9,10),dt.date(2024,9,12))
    self.assertEqual(max((r['wind'] for r in rows if r['wind'] is not None),default=None),minute)
    rows=iem.parse(z.read(sid+'-all-metar.csv').decode(),sid,dt.date(2024,9,10),dt.date(2024,9,12))
    self.assertEqual(max(r['wind'] for r in rows if r['wind'] is not None),aviation)
    if sid=='KBTR':
     peak=max(rows,key=lambda r:r['wind'] or 0)
     self.assertEqual(peak['time'],dt.datetime(2024,9,12,1,tzinfo=UTC))
     self.assertEqual(peak['report_type'],1)
     self.assertEqual(peak['wind_averaging_period_minutes'],2)
     self.assertIn('03027G35KT',peak['raw']['metar'])
 def test_five_minute_rain_not_added_to_hourly_total(self):
  with zipfile.ZipFile(EVIDENCE) as z:
   rows=iem.parse(z.read('KBTR-all-metar.csv').decode(),'KBTR',dt.date(2024,9,10),dt.date(2024,9,12))
  hf=[r for r in rows if r.get('report_type')==1]
  self.assertTrue(hf)
  self.assertIsNone(iem.rain_total(hf,dt.datetime(2024,9,12,tzinfo=UTC),dt.datetime(2024,9,12,2,tzinfo=UTC)))
 def test_cross_source_precedence_and_provenance(self):
  wb=openpyxl.Workbook();wb.active.title='Wind and Pressure';wb.create_sheet('Rainfall')
  audit=Audit();start=dt.date(2024,9,10);end=dt.date(2024,9,12)
  inventory=[{'id':s,'row':r,'network':'ASOS'} for r,s in enumerate(['KBTR','KMSY','KNEW','KPQL'],2)]
  for st,v in zip(inventory,[24,50,42,23]):
   wb['Wind and Pressure'].cell(st['row'],11).value=v
   audit.add('Wind and Pressure',st['row'],st['id'],'wind',v,'kn',dt.datetime(2024,9,12,tzinfo=UTC),'https://example.org/')
  def collect(s,*args):
   with zipfile.ZipFile(EVIDENCE) as z:rows=iem.parse(z.read('K'+s+'-all-metar.csv').decode(),s,start,end)
   return rows,'https://mesonet.agron.iastate.edu/cgi-bin/request/asos.py'
  with patch.object(iem,'collect',side_effect=collect),patch.object(iem,'inventory',side_effect=lambda w,t:inventory if t=='Wind and Pressure' else []):
   iem.populate(wb,[],start,end,Counter(),audit)
  self.assertEqual([wb['Wind and Pressure'].cell(r,11).value for r in range(2,6)],[27,50,42,23])
  winds=[e for e in audit.entries if e['variable']=='wind']
  self.assertEqual(len(winds),4)
  btr=next(e for e in winds if e['site_id']=='KBTR')
  self.assertEqual(btr['source_kind'],'ASOS five-minute report')
  self.assertEqual(btr['averaging_period_minutes'],2)
  self.assertIn('03027G35KT',btr['source_record']['metar'])
  self.assertEqual(len(btr['aviation_wind_comparison']),2)
  self.assertTrue(all(e['selection_policy'] for e in winds))
 def test_request_includes_hf_without_nonstandard_type_two(self):
  class Response:
   text='station,valid,sknt\n';url='https://example.org/'
   def raise_for_status(self):pass
  with patch.object(iem.requests,'get',return_value=Response()) as get:
   iem.collect('KBTR',dt.date(2024,9,10),dt.date(2024,9,12))
  self.assertEqual(get.call_args.kwargs['params']['report_type'],[1,3,4])
