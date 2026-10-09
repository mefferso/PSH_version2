import datetime as dt
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import build
import export_dashboard
import validate
from common import UTC,inventory
from copy import copy
ROOT=Path(__file__).resolve().parents[1]
class PipelineTests(unittest.TestCase):
    def test_real_workbook_preservation_provenance_summary_and_dashboard(self):
        def fake(w,q,a,b,c,audit):
            from iem import write_wind
            rows=[{'wind':40.0,'gust':60.0,'pressure':990.0,'dir':90,'time':dt.datetime(2024,9,11,12,tzinfo=UTC)}]
            write_wind(w['Wind and Pressure'],2,rows,'KBTR',audit,'https://mesonet.agron.iastate.edu/','controlled integration fixture')
            q.append(['Wind and Pressure','KBTR','ASOS','REVIEW REQUIRED','fixture','https://mesonet.agron.iastate.edu/'])
        with tempfile.TemporaryDirectory() as d,patch.dict(os.environ,{'STORM_NAME':'Hurricane Francine','START_UTC':'2024-09-11','END_UTC':'2024-09-11','PSH_OFFLINE':'1'},clear=True):
            out=Path(d)/'output';site=Path(d)/'site'
            with patch.object(build,'OUT',out),patch.object(build,'TEMPLATE',ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'),patch.object(build,'collectors',return_value=[fake]):
                manifest=build.build()
            self.assertEqual(manifest['measurement_count'],3)
            self.assertTrue(validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'))
            generated=load_workbook(out/manifest['workbook']);original=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
            self.assertEqual(generated['Summary']['D20'].value,40)
            self.assertNotIn('[',generated['Rainfall']['B230'].value)
            self.assertNotIn('1-hour',generated['Water Level']['B117'].value)
            for tab,cols in [('Wind and Pressure',10),('Rainfall',7),('Water Level',6)]:
                for r in [1]+[st["row"] for st in inventory(original,tab)]:
                    for c in range(1,cols+1):
                        x=original[tab].cell(r,c);y=generated[tab].cell(r,c)
                        self.assertEqual(x.value,y.value)
                        for attr in ("font","fill","border","alignment","protection","number_format"):
                            self.assertEqual(copy(getattr(x,attr)),copy(getattr(y,attr)))
                        self.assertEqual(x.hyperlink.target if x.hyperlink else None,y.hyperlink.target if y.hyperlink else None)
            payload=export_dashboard.export(out,site)
            self.assertEqual(set(payload['tabs']),set(generated.sheetnames)|{'Provenance'})
            self.assertTrue(payload['tabs']['Summary']['rows'])
            self.assertTrue((site/'review-outputs.zip').exists())
            self.assertEqual(len(payload['audit']),3)
    def test_validator_rejects_measurement_without_provenance(self):
        with tempfile.TemporaryDirectory() as d:
            out=Path(d);w=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx');w.save(out/'test.xlsx')
            (out/'QC.json').write_text(json.dumps({'workbook':'test.xlsx','measurement_count':0}))
            (out/'provenance.json').write_text('[]')
            with self.assertRaises(ValueError):validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')

class StaleNarrativeTests(unittest.TestCase):
    def test_example_flooding_and_impact_values_are_cleared(self):
        w=load_workbook(ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
        build.reset(w)
        self.assertIsNone(w['Inland Flooding']['D6'].value)
        self.assertIsNone(w['Inland Flooding']['D7'].value)
        self.assertIsNone(w['Inland Flooding']['B3'].value)
        self.assertEqual(w['Inland Flooding']['A6'].value,'Harrison')

class PublicationGateTests(unittest.TestCase):
    def test_rain_interval_summary_csv_and_archive_tampering_are_rejected(self):
        def fixture(w,q,start,end,counts,audit):
            from cocorahs import rain_bounds
            a,b=rain_bounds(start,end)
            w['Rainfall'].cell(98,8).value=5
            audit.add('Rainfall',98,'LA-AS-02','rain',5,'in',b,'https://www.cocorahs.org/',interval_start=a)
        with tempfile.TemporaryDirectory() as d,patch.dict(os.environ,{'STORM_NAME':'Hurricane Francine','START_UTC':'2024-09-11','END_UTC':'2024-09-11','PSH_OFFLINE':'1'},clear=True):
            out=Path(d)/'output';site=Path(d)/'site'
            with patch.object(build,'OUT',out),patch.object(build,'TEMPLATE',ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx'),patch.object(build,'collectors',return_value=[fixture]):
                meta=build.build()
            validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
            audit=json.loads((out/'provenance.json').read_text());original=(out/'provenance.json').read_text()
            audit[0]['interval_start_utc']='1999-01-01T00:00:00Z';(out/'provenance.json').write_text(json.dumps(audit))
            with self.assertRaisesRegex(ValueError,'Rain audit interval'):validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
            (out/'provenance.json').write_text(original)
            path=out/meta['workbook'];saved=path.read_bytes();w=load_workbook(path);w['Summary']['D76']=9999;w.save(path)
            with self.assertRaisesRegex(ValueError,'Summary differs'):validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
            path.write_bytes(saved)
            (out/'unrelated_old_storm.xlsx').write_bytes(b'old')
            export_dashboard.export(out,site);validate.validate_site(out,site)
            import zipfile
            with zipfile.ZipFile(site/'review-outputs.zip') as z:self.assertNotIn('unrelated_old_storm.xlsx',z.namelist())
            with (out/'csv/Rainfall_REVIEW.csv').open('a') as f:f.write('bad\n')
            with self.assertRaisesRegex(ValueError,'CSV differs'):validate.validate(out,ROOT/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
