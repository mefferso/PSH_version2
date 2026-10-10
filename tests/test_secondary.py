import datetime as dt
import sys
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from common import UTC
import cocorahs
import weatherstem

class SecondaryTests(unittest.TestCase):
    def test_cocorahs_exact_id_padding_and_complete_periods(self):
        a=dt.datetime(2024,9,10,12,tzinfo=UTC);b=a+dt.timedelta(days=2)
        reports=[{'stationNumber':'LA-ST-11','obsDateTime':f'2024-09-{day}T12:00:00Z','gaugeCatch':v,'numDays':1} for day,v in [(11,4.5),(12,5.13)]]
        self.assertAlmostEqual(cocorahs.total(reports,'LA-ST-011',a,b),9.63)
        self.assertIsNone(cocorahs.total(reports,'LA-JF-20',a,b))
        self.assertIsNone(cocorahs.total(reports[:1],'LA-ST-11',a,b))
        reports[1]['gaugeCatchIsTrace']=True
        self.assertIsNone(cocorahs.total(reports,'LA-ST-11',a,b))

    def test_weatherstem_units_and_mean_period_required(self):
        metadata={'id':123,'transmitters':[{'sensors':[{'id':1,'name':'Anemometer','unit':'mph'},
             {'id':2,'name':'10 Minute Wind Gust','unit':'mph'},{'id':3,'name':'Barometer','unit':'inHg'}]}]}
        raw=[['Timestamp','Anemometer','10 Minute Wind Gust','Barometer'],['2024-09-11 12:00',40,50,29.5]]
        rows=weatherstem.parse(raw,metadata,dt.date(2024,9,11),dt.date(2024,9,11))
        self.assertIsNone(rows[0]['wind']) # no validated sustained averaging period
        self.assertAlmostEqual(rows[0]['gust'],50*0.8689762419)
        self.assertIsNone(rows[0]['pressure']) # barometer not proven MSLP
        metadata['transmitters'][0]['sensors'][0]['averaging_period_minutes']=2
        self.assertAlmostEqual(weatherstem.parse(raw,metadata,dt.date(2024,9,11),dt.date(2024,9,11))[0]['wind'],40*0.8689762419)

    def test_weatherstem_station_link_cannot_choose_nearby_station(self):
        self.assertEqual(weatherstem.link_parts('https://eastbatonrouge.weatherstem.com/data?refer=/alexbox'),('eastbatonrouge','alexbox'))
        with self.assertRaises(ValueError):weatherstem.link_parts('https://example.org/data?refer=/alexbox')

class CapturedWeatherSTEMTests(unittest.TestCase):
    def test_live_metadata_symbol_and_seconds_timestamp(self):
        meta={'id':1313,'transmitters':[{'sensors':[{'id':27959,'name':'10 Minute Wind Gust','units':{'name':'Miles Per Hour','symbol':'mph'}}]}]}
        raw=[['Record ID','Timestamp','10 Minute Wind Gust'],[7318110,'2024-09-11 12:01:39',35]]
        rows=weatherstem.parse(raw,meta,dt.date(2024,9,11),dt.date(2024,9,11))
        self.assertEqual(len(rows),1)
        self.assertAlmostEqual(rows[0]['gust'],30.4141684665)
        self.assertEqual(rows[0]['time'].second,39)

class PeakDirectionSafetyTests(unittest.TestCase):
    def test_concurrent_wind_vane_does_not_prove_peak_gust_direction(self):
        from openpyxl import Workbook
        from common import Audit
        from iem import write_wind
        sheet=Workbook().active
        t=dt.datetime(2024,9,11,12,tzinfo=UTC)
        rows=[{'time':t,'wind':30,'gust':52,'dir':180,'gust_dir':None,'pressure':None}]
        audit=Audit()
        self.assertEqual(write_wind(sheet,2,rows,'TEST',audit,'https://example.org/','fixture'),2)
        self.assertEqual(sheet.cell(2,12).value,180)
        self.assertIsNone(sheet.cell(2,18).value)
        rows[0]['gust_dir']=220
        other=Workbook().active
        write_wind(other,2,rows,'TEST',Audit(),'https://example.org/','fixture')
        self.assertEqual(other.cell(2,18).value,220)
