import datetime as dt
import sys
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import coastal_water
from common import UTC

class ElapsedCoverageTests(unittest.TestCase):
 def test_future_is_not_missing(self):
  a=dt.datetime(2026,10,8,tzinfo=UTC);b=a+dt.timedelta(hours=4)
  rows=[dict(time_utc=(a+dt.timedelta(hours=i)).isoformat(),original_value=i,original_unit='ft',quality_code=0,eligible=True,value_ft=i) for i in range(2)]
  c=coastal_water.summarize(rows,a,b,3600,as_of=a+dt.timedelta(hours=2))
  self.assertEqual(c['expected_count'],2);self.assertEqual(c['missing_count'],0)
  self.assertTrue(c['provisional']);self.assertTrue(c['complete']);self.assertEqual(c['peak']['value_ft'],1)
 def test_before_and_after(self):
  a=dt.datetime(2026,10,8,tzinfo=UTC);b=a+dt.timedelta(hours=4)
  c=coastal_water.summarize([],a,b,3600,as_of=a-dt.timedelta(hours=1))
  self.assertEqual(c['expected_count'],0);self.assertEqual(c['missing_count'],0);self.assertTrue(c['provisional'])
  c=coastal_water.summarize([],a,b,3600,as_of=b+dt.timedelta(hours=1))
  self.assertEqual(c['missing_count'],4);self.assertFalse(c['provisional'])
 def test_hml_cadence_not_primary_cadence(self):
  a=dt.datetime(2026,10,8,tzinfo=UTC)
  rows=[{'time_utc':(a+dt.timedelta(minutes=15*i)).isoformat()} for i in range(8)]
  self.assertEqual(coastal_water.observed_cadence(rows),900)
  self.assertIsNone(coastal_water.observed_cadence(rows[:2]))

class MarineRankingTests(unittest.TestCase):
 def test_all_networks_and_independent_gusts(self):
  from openpyxl import load_workbook
  import products,observation_review
  w=load_workbook(Path(__file__).resolve().parents[1]/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx');sheet=w['Wind and Pressure']
  from common import inventory
  for st in inventory(w,'Wind and Pressure'):
   for col in (11,17):sheet.cell(st['row'],col).value=None
  for row,network,wind,gust,height in [(2,'ASOS',40,None,10),(3,'AWOS',45,50,19),(4,'CMAN',30,60,4),(5,'WeatherFlow',80,90,None),(6,'WLON',70,100,20)]:
   for col,value in [(7,'M'),(8,network),(9,height),(10,2),(11,wind),(17,gust)]:sheet.cell(row,col).value=value
  products.summaries(w)
  self.assertEqual([w['Summary'].cell(r,3).value for r in range(48,51)],[45,40,30])
  self.assertEqual([w['Summary'].cell(r,3).value for r in range(62,64)],[60,50])
  rows={x['row']:x for x in observation_review.marine(w)}
  self.assertIn('unresolved',rows[5]['wind_reason']);self.assertIn('>=20',rows[6]['gust_reason'])
 def test_unknown_mean_excludes_wind_only(self):
  from openpyxl import load_workbook
  import observation_review
  w=load_workbook(Path(__file__).resolve().parents[1]/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx');s=w['Wind and Pressure']
  for col,v in [(7,'M'),(9,10),(10,None),(11,30),(17,40)]:s.cell(2,col).value=v
  r=next(x for x in observation_review.marine(w) if x['row']==2)
  self.assertFalse(r['wind_eligible']);self.assertTrue(r['gust_eligible'])

class RainElapsedTests(unittest.TestCase):
 def test_future_hours_and_unfinished_hour_not_missing(self):
  import os
  from unittest.mock import patch
  import hourlyprecip
  a=dt.datetime(2026,10,8,tzinfo=UTC);b=a+dt.timedelta(hours=4)
  body='station,network,valid,precip_in\nMSY,LA_ASOS,2026-10-08 00:00,0.1\nMSY,LA_ASOS,2026-10-08 01:00,0.2\nMSY,LA_ASOS,2026-10-08 02:00,9\n'
  with patch.dict(os.environ,PSH_AS_OF_UTC='2026-10-08T02:30:00Z'):
   info=hourlyprecip.coverage_details(body,'MSY','LA_ASOS',a,b)
   self.assertEqual(info['expected_hours'],2);self.assertEqual(info['missing_hours_utc'],[])
   self.assertEqual(info['observed_sum_inches'],0.3);self.assertTrue(info['provisional'])
 def test_partial_intervals_cannot_include_future_reports(self):
  import os
  from unittest.mock import patch
  from cocorahs import observed_partial_sum
  a=dt.datetime(2026,10,8,tzinfo=UTC)
  with patch.dict(os.environ,PSH_AS_OF_UTC='2026-10-08T02:00:00Z'):
   r=observed_partial_sum([{'start':a,'end':a+dt.timedelta(hours=1),'value':0.1},{'start':a+dt.timedelta(hours=1),'end':a+dt.timedelta(hours=3),'value':10}],a,a+dt.timedelta(hours=4))
   self.assertEqual(r[0],0.1)

class SynopticNativeRainTests(unittest.TestCase):
 def test_native_sensor_key_and_no_parallel_sum(self):
  import synoptic
  a=dt.datetime(2024,9,10,tzinfo=UTC);b=a+dt.timedelta(days=1)
  reports=[dict(first_report=a.isoformat(),last_report=b.isoformat(),count=24,interval=1,total=1.2)]
  site={'STID':'TEST','OBSERVATIONS':{'precip_accum_one_hour':reports}}
  p={'SUMMARY':{'RESPONSE_CODE':1},'UNITS':{'precipitation':'Inches'},'STATION':[site]}
  self.assertEqual(synoptic.precip_total(p,'TEST',a,b),1.2)
  site['OBSERVATIONS']['precip_accum_24_hour']=reports
  self.assertIsNone(synoptic.precip_total(p,'TEST',a,b))
 def test_native_off_boundary_not_storm_total(self):
  import synoptic
  a=dt.datetime(2024,9,10,12,tzinfo=UTC);b=a+dt.timedelta(days=1)
  site={'STID':'TEST','OBSERVATIONS':{'precip_accum_one_hour':[dict(first_report=(a-dt.timedelta(minutes=7)).isoformat(),last_report=(b-dt.timedelta(minutes=7)).isoformat(),count=24,interval=1,total=1.2)]}}
  p={'SUMMARY':{'RESPONSE_CODE':1},'UNITS':{'precipitation':'Inches'},'STATION':[site]}
  self.assertIsNone(synoptic.precip_total(p,'TEST',a,b));self.assertIsNone(synoptic.partial_precip_from_reports(p,'TEST',a,b))

class OperationalRegressionTests(unittest.TestCase):
 def test_actual_run116_rankings_match_eligible_workbook_rows(self):
  import json,products
  from openpyxl import load_workbook
  from common import inventory
  root=Path(__file__).resolve().parents[1]
  w=load_workbook(root/'Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx');sheet=w['Wind and Pressure']
  for st in inventory(w,'Wind and Pressure'):
   for c in (11,17):sheet.cell(st['row'],c).value=None
  data=json.loads((root/'tests/fixtures/isaias_run116_marine.json').read_text())
  for row in data['rows']:
   for col,value in enumerate(row['values'],1):sheet.cell(row['row'],col).value=value
  products.summaries(w)
  self.assertEqual([w['Summary'].cell(r,3).value for r in range(48,54)],[36.9,31.1,29,28,27,25.1])
  self.assertEqual([w['Summary'].cell(r,3).value for r in range(62,68)],[52.1,39.1,38.1,36,36,35])
  self.assertIsNone(w['Summary'].cell(54,3).value)
 def test_new_direct_datum_evidence_does_not_backdate_francine(self):
  mapping=coastal_water.load_mapping()
  for sid in ('SBEL1','PRSL1'):
   m=mapping[sid];loc={'description':m['description'],'vertical-datum':m['location_vertical_datum']}
   self.assertTrue(coastal_water.qualify(m,loc,dt.datetime(2026,10,8,tzinfo=UTC),dt.datetime(2026,10,11,tzinfo=UTC))[0])
   self.assertFalse(coastal_water.qualify(m,loc,dt.datetime(2024,9,10,tzinfo=UTC),dt.datetime(2024,9,13,tzinfo=UTC))[0])
   loc['description']='Gage relocated'
   self.assertFalse(coastal_water.qualify(m,loc,dt.datetime(2026,10,8,tzinfo=UTC),dt.datetime(2026,10,11,tzinfo=UTC))[0])

class RainReportingCadenceTests(unittest.TestCase):
 def test_daily_report_not_due_is_pending(self):
  import observation_review
  a=dt.datetime(2026,10,8,12,tzinfo=UTC);b=a+dt.timedelta(days=3)
  r=observation_review.interval_coverage([{'start':a.isoformat(),'end':(a+dt.timedelta(days=1)).isoformat()}],a,b,a+dt.timedelta(hours=42))
  self.assertEqual(r['missing_due_report_count'],0);self.assertEqual(r['pending_elapsed_hours'],18);self.assertEqual(r['future_hours'],30)
 def test_elapsed_missing_day_is_retained_as_gap(self):
  import observation_review
  a=dt.datetime(2024,9,10,12,tzinfo=UTC);b=a+dt.timedelta(days=3)
  r=observation_review.interval_coverage([{'start':a.isoformat(),'end':(a+dt.timedelta(days=1)).isoformat()}],a,b,b)
  self.assertEqual(r['missing_due_report_count'],2);self.assertEqual(r['observed_hours'],24)

class ArchivedCoCoRaHSTests(unittest.TestCase):
 def test_no_clock_no_storm_total_and_exact_identity(self):
  import rain_archive_review
  p={'data':[{'station':'LA-JF-5','id':'LA-JF-5','date':'2026-10-09','precip':.01,'name':'River Ridge 0.7 N'},{'station':'LA-JF-6','id':'LA-JF-6','date':'2026-10-09','precip':99}]}
  r=rain_archive_review.parse(p,'LA-JF-05','LA_COCORAHS',dt.date(2026,10,8),dt.date(2026,10,10))
  self.assertEqual(len(r),1);self.assertEqual(r[0]['reported_inches'],.01);self.assertFalse(r[0]['can_populate_psh']);self.assertIsNone(r[0]['actual_observation_time_utc'])
 def test_current_report_date_and_missing_value_not_used(self):
  import rain_archive_review
  p={'data':[{'station':'LA-JF-5','id':'LA-JF-5','date':'2026-10-10','precip':9}]}
  self.assertEqual(rain_archive_review.parse(p,'LA-JF-5','LA_COCORAHS',dt.date(2026,10,8),dt.date(2026,10,10)),[])

class RegressionBoundaryTests(unittest.TestCase):
 def test_decimal_tolerance_boundary_is_inclusive(self):
  import regression
  from openpyxl import Workbook
  from copy import deepcopy
  w=Workbook();w.active.title='Rainfall';w.create_sheet('Wind and Pressure');w.create_sheet('Water Level')
  s=w['Rainfall']
  for col,v in [(1,'TEST'),(2,'Test'),(3,30),(4,-90),(7,'COOP'),(8,5.55)]:s.cell(2,col).value=v
  s.cell(3,1).value='Rainfall Start Time';s.cell(3,2).value='1200 UTC Sep 10 2024'
  s.cell(4,1).value='Rainfall End Time';s.cell(4,2).value='1200 UTC Sep 12 2024'
  g=deepcopy(w);g['Rainfall'].cell(2,8).value=5.43
  self.assertEqual(regression.compare(g,w)['within_tolerance'],1)
  g['Rainfall'].cell(2,8).value=5.42
  self.assertEqual(regression.compare(g,w)['discrepancies'],1)
