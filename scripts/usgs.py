"""Conservative USGS NAVD88 elevation collector via modern Water Data OGC API.

Only direct NAVD88 elevation parameter codes (63160, 62620 and 62615, feet) are written.
Stage (00065) is deliberately not converted without validated vertical datum metadata.
"""
import datetime as dt
import csv
from pathlib import Path
import math
import os
import re
import time
from urllib.parse import urlparse
import requests
from openpyxl.cell.cell import MergedCell
from common import inventory

BASE = "https://api.waterdata.usgs.gov/ogcapi/v1/collections/continuous/items"
PARAMETERS = ("63160", "62620", "62615")
USGS_DOMAINS = {"waterdata.usgs.gov", "www.waterdata.usgs.gov",
                "nwis.waterdata.usgs.gov", "water.usgs.gov", "www.water.usgs.gov"}

def site_from_link(cell):
    url = urlparse(cell.hyperlink.target if cell.hyperlink else "")
    if url.hostname not in USGS_DOMAINS:
        return None
    candidates = re.findall(r"(?:monitoring-location/|site_no=|sites=|site/)([0-9]{8,15})", url.geturl())
    return candidates[0] if candidates else None

def parse_observations(payload, site, start, end):
    found = []
    for feature in payload.get("features", []):
        p = feature.get("properties") or {}
        if p.get("monitoring_location_id") != "USGS-" + site:
            continue
        code = str(p.get("parameter_code") or "")
        if code not in PARAMETERS:
            continue
        unit = str(p.get("unit_of_measure") or "").lower()
        if unit not in ("ft", "feet", "foot"):
            continue
        try:
            value = float(p["value"])
            moment = dt.datetime.fromisoformat(str(p["time"]).replace("Z", "+00:00"))
            if moment.tzinfo is None:
                continue
            moment = moment.astimezone(dt.timezone.utc)
        except (ValueError, TypeError, KeyError, OverflowError):
            continue
        if math.isfinite(value) and -30 < value < 50 and start <= moment.date() <= end:
            found.append((value, moment, p))
    return found

def verify_parameter(payload, code):
    p=payload.get("properties") or {}
    if str(p.get("id")) != code or p.get("unit_of_measure") != "ft" or not re.search(r"NAVD\s*(?:1988|88)",p.get("parameter_description", ""), re.I):
        raise ValueError("Parameter metadata does not establish NAVD88 feet")
    return True

def request_headers():
    key=os.environ.get('USGS_API_KEY')
    return {'X-Api-Key':key} if key else {}

def parameter_evidence(code, session=requests):
    url="https://api.waterdata.usgs.gov/ogcapi/v1/collections/parameter-codes/items/"+code
    r=session.get(url,params={"f":"json"},timeout=20,headers=request_headers());r.raise_for_status()
    verify_parameter(r.json(),code)
    return r.url

def collect(site, start, end, session=requests):
    # parameter_code is an exact string filter: comma-separated codes return no series.
    urls=[]
    for code in PARAMETERS:
        params={"f":"json","monitoring_location_id":"USGS-"+site,"parameter_code":code,
            "time":f"{start.isoformat()}T00:00:00Z/{end.isoformat()}T23:59:59Z","limit":10000,"skipGeometry":"true"}
        readings=[];next_url=BASE
        for page in range(12):
            r=session.get(next_url,params=params if page==0 else None,timeout=20,headers=request_headers())
            r.raise_for_status();data=r.json()
            if data.get("type")!="FeatureCollection":raise ValueError("Unexpected USGS API response")
            urls.append(r.url)
            for value,moment,meta in parse_observations(data,site,start,end):
                readings.append((value,moment,dict(meta,source_url=r.url)))
            nexts=[x.get("href") for x in data.get("links",[]) if x.get("rel")=="next"]
            if not nexts:break
            next_url=nexts[0]
            if not next_url.startswith("https://api.waterdata.usgs.gov/"):raise ValueError("Unsafe pagination URL")
        else:raise ValueError("USGS pagination limit reached; incomplete series discarded")
        series={(x[2].get("time_series_id"),x[2].get("parameter_code")) for x in readings}
        if len(series)>1:raise ValueError("Multiple USGS series/parameters; explicit series selection required")
        if readings:return max(readings,key=lambda x:x[0]),urls
    return None,urls

def parse_stage_candidates(payload,site,start,end):
    """Reported stage (00065) in feet; NEVER a NAVD88 elevation."""
    values=[]
    for feature in payload.get('features',[]):
        p=feature.get('properties') or {}
        if p.get('monitoring_location_id')!='USGS-'+site:continue
        if str(p.get('parameter_code'))!='00065':continue
        if str(p.get('unit_of_measure') or '').lower() not in ('ft','feet','foot'):continue
        try:
            value=float(p['value'])
            moment=dt.datetime.fromisoformat(str(p['time']).replace('Z','+00:00'))
            if moment.tzinfo is None or not math.isfinite(value):continue
            moment=moment.astimezone(dt.timezone.utc)
        except (ValueError,TypeError,KeyError,OverflowError):continue
        if -100<value<200 and start<=moment.date()<=end:
            values.append((value,moment))
    return values

def stage_peak(site,start,end,session=requests):
    params={'f':'json','monitoring_location_id':'USGS-'+site,
        'parameter_code':'00065',
        'time':f'{start.isoformat()}T00:00:00Z/{end.isoformat()}T23:59:59Z',
        'limit':10000,'skipGeometry':'true'}
    readings=[];url=BASE
    for page in range(12):
        r=session.get(url,params=params if page==0 else None,timeout=20,headers=request_headers())
        r.raise_for_status();payload=r.json()
        if payload.get('type')!='FeatureCollection':raise ValueError('Unexpected USGS stage response')
        readings.extend(parse_stage_candidates(payload,site,start,end))
        nexts=[x.get('href') for x in payload.get('links',[]) if x.get('rel')=='next']
        if not nexts:break
        url=nexts[0]
        if not url.startswith('https://api.waterdata.usgs.gov/'):
            raise ValueError('Unsafe stage pagination URL')
    else:raise ValueError('USGS stage pagination incomplete')
    return max(readings,key=lambda x:x[0]) if readings else None

def populate(wb, qc, start, end, counts, audit=None, output_dir=None):
    from concurrent.futures import ThreadPoolExecutor
    sheet = wb["Water Level"]
    ids={site_from_link(sheet.cell(st["row"],1)) for st in inventory(wb,"Water Level") if st["network"].upper()=="USGS"}
    ids.discard(None)
    def fetch(site):
        try:return site,collect(site,start,end)
        except (requests.RequestException,ValueError) as e:return site,e
    with ThreadPoolExecutor(max_workers=3) as pool:cache=dict(pool.map(fetch,ids))
    # Request stage-only review values concurrently to avoid serial network
    # delays for dozens of stations lacking direct vertical elevations.
    stage_ids=[sid for sid,result in cache.items()
               if not isinstance(result,Exception) and result[0] is None]
    def fetch_stage(sid):
        try:return sid,stage_peak(sid,start,end)
        except (requests.RequestException,ValueError) as exc:return sid,exc
    with ThreadPoolExecutor(max_workers=6) as pool:
        stage_cache=dict(pool.map(fetch_stage,stage_ids))
    stage_review=[]
    evidence={}
    for code in PARAMETERS:
        try:evidence[code]=parameter_evidence(code)
        except (requests.RequestException,ValueError):evidence[code]=None
    for row in range(2, sheet.max_row + 1):
        from common import identifier
        site = identifier(sheet.cell(row, 1).value)
        source = str(sheet.cell(row, 13).value or "").strip().upper()
        usgs = site_from_link(sheet.cell(row, 1))
        if not usgs or not site:
            continue
        if source != "USGS":
            qc.append(["Water Level", site, source, "CHECK SOURCE",
                       "USGS hyperlink exists, but source column M is not USGS", ""])
            continue
        # No suitable direct-elevation series: leave measurement blank, preserve station metadata.
        for col in (7,8,9,10,11,12,14,15):
            cell = sheet.cell(row,col)
            if not isinstance(cell,MergedCell):
                cell.value = None
        try:
            result=cache[usgs]
            if isinstance(result,Exception):raise result
            peak,urls=result
            if not peak:
                counts["usgs_no_direct_navd88"] += 1
                detail="Direct NAVD88 ft codes 63160/62620/62615 unavailable; 00065 stage was NOT converted"
                candidate=stage_cache.get(usgs)
                if isinstance(candidate,Exception):
                    detail+="; stage review lookup unavailable: "+type(candidate).__name__
                elif candidate:
                    stage,when=candidate
                    detail+=f"; MANUAL REVIEW ONLY: observed 00065 stage peak {stage:.2f} ft at {when.isoformat()} (gage height, NOT NAVD88 or inundation depth)"
                    counts["usgs_stage_review_candidates"]+=1
                    stage_review.append({'site_id':site,'usgs_site_number':usgs,
                        'peak_gage_height_ft':round(stage,2),'peak_time_utc':when.isoformat(),
                        'datum':'GAGE HEIGHT — NOT NAVD88',
                        'qualification':'MANUAL DATUM REVIEW REQUIRED',
                        'source_url':'https://waterdata.usgs.gov/monitoring-location/'+usgs+'/'})
                qc.append(["Water Level",site,"USGS","NO DIRECT NAVD88",detail,
                           urls[0] if urls else ""])
                continue
            value, moment, meta = peak
            citation=evidence.get(meta.get("parameter_code"))
            if not citation:raise ValueError("NAVD88 parameter metadata unavailable")
            if audit:
                audit.add("Water Level",row,site,"water",round(value,2),"ft",moment,
                          meta["source_url"],datum="NAVD88",evidence=citation,
                          raw_value=value,details="Direct parameter "+str(meta.get("parameter_code"))+"; provisional/parameter metadata review required")
            sheet.cell(row,7).value = round(value,2)
            sheet.cell(row,8).value = "NAVD88"
            sheet.cell(row,9).value = moment.strftime("%H%M")
            sheet.cell(row,10).value = moment.day
            sheet.cell(row,11).value = moment.month
            sheet.cell(row,12).value = moment.year
            sheet.cell(row,14).value = "I"
            sheet.cell(row,15).value = ("USGS preliminary review required; "+
                                        "direct NAVD88 elevation parameter "+str(meta.get("parameter_code")))
            counts["usgs_navd88_collected"] += 1
            qc.append(["Water Level",site,"USGS","REVIEW REQUIRED",
                       f"{value:.2f} ft NAVD88 UTC {moment:%Y-%m-%d %H:%M}; parameter {meta.get('parameter_code')}",
                       urls[0] if urls else ""])
        except Exception as exc:
            counts["usgs_errors"] += 1
            response=getattr(exc,"response",None)
            code=getattr(response,"status_code",None)
            status="RATE LIMITED" if code==429 else "ERROR"
            detail="USGS public/API-key quota exceeded; retry after reset or configure USGS_API_KEY" if code==429 else type(exc).__name__+": direct NAVD88 request unsuccessful"
            qc.append(["Water Level",site,"USGS",status,detail,
                       "https://waterdata.usgs.gov/monitoring-location/"+usgs+"/"])
        time.sleep(0.1)

    path=(Path(output_dir) if output_dir is not None else Path('output'))/'USGS_stage_review.csv'
    path.parent.mkdir(parents=True,exist_ok=True)
    columns=['site_id','usgs_site_number','peak_gage_height_ft','peak_time_utc',
             'datum','qualification','source_url']
    with path.open('w',newline='',encoding='utf-8') as handle:
        writer=csv.DictWriter(handle,fieldnames=columns)
        writer.writeheader()
        writer.writerows(stage_review)
