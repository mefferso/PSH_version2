"""Reproducible comparison to a saved operational Pages manifest."""
import argparse
import datetime as dt
import json
from pathlib import Path
from openpyxl import load_workbook
from common import inventory,finite,timestamp
from observation_review import marine

def audit(out,baseline):
    out=Path(out);meta=json.loads((out/'QC.json').read_text());book=load_workbook(out/meta['workbook'])
    old=json.loads(Path(baseline).read_text());rain=[]
    previous={str(r['v'][0]):r['v'][7] for r in old['tabs']['Rainfall']['rows'] if len(r['v'])>7 and isinstance(r['v'][7],(int,float))}
    for st in inventory(book,'Rainfall'):
        value=finite(book['Rainfall'].cell(st['row'],8).value,0,100)
        if value is not None:rain.append(dict(site_id=st['id'],network=st['network'],row=st['row'],before=previous.get(st['id']),after=value,flag=book['Rainfall'].cell(st['row'],9).value))
    coastal=json.loads((out/'coastal_water_audit.json').read_text());raw=json.loads((out/'coastal_water_observations.json').read_text());checks=[]
    for r in coastal:
        source=raw[r['site_id']];times={}
        # Recompute from retained raw responses, independent of collector rows.
        for page in source['pages']:
            p=page.get('payload')
            if not p:continue
            factor={'ft':1,'m':1/.3048,'cm':1/30.48}[p['units']]
            cols={x['name']:x['ordinal']-1 for x in p['value-columns']}
            for v in p['values']:
                t=dt.datetime.fromtimestamp(v[cols['date-time']]/1000,dt.timezone.utc)
                if not timestamp(r['requested_start_utc'])<=t<timestamp(r['elapsed_end_utc']):continue
                value=finite(v[cols['value']],-100/factor,100/factor)
                q=v[cols['quality-code']]
                if value is None or q not in (0,3):continue
                times.setdefault(t,set()).add(value*factor)
        observations=sorted((t,next(iter(v))) for t,v in times.items() if len(v)==1)
        peak=max(observations,key=lambda p:p[1]) if observations else None
        if source.get('source_kind','').startswith('NWS HML'):
            import csv,io
            rows=[]
            for page in source['pages']:
                reader=csv.DictReader(io.StringIO(page['payload_csv']))
                for v in reader:
                    t=dt.datetime.strptime(v['valid[UTC]'],'%Y-%m-%d %H:%M').replace(tzinfo=dt.timezone.utc);value=finite(v.get('Stage[ft]'))
                    if t and value is not None and timestamp(r['requested_start_utc'])<=t<timestamp(r['elapsed_end_utc']):rows.append((t,value))
            peak=max(sorted(rows),key=lambda p:p[1]) if rows else None
        matches=(r['peak_ft'] is None and peak is None) or (peak is not None and abs(peak[1]-r['peak_ft'])<1e-9 and peak[0]==timestamp(r['peak_time_utc']))
        if not matches:raise ValueError('Raw source peak mismatch '+r['site_id'])
        checks.append(dict(site_id=r['site_id'],peak_ft=r['peak_ft'],time_utc=r['peak_time_utc'],independent_raw_peak_matches=True))
    rankings=marine(book)
    return dict(baseline_workflow=old['meta']['workflow_url'],baseline_generated=old['meta']['generated_at_utc'],as_of_utc=meta.get('as_of_utc'),
                observation_start_utc=meta['observation_start_utc'],observation_end_utc=meta['observation_end_utc'],
                coastal_before={'observed':27,'qualified':11},coastal_after={'observed':sum(r['peak_ft'] is not None for r in coastal),'qualified':sum(r['can_populate_psh'] for r in coastal)},
                datum_review=[r['site_id'] for r in coastal if r['peak_ft'] is not None and not r['can_populate_psh']],
                no_observations=[r['site_id'] for r in coastal if r['peak_ft'] is None],coastal_raw_checks=checks,
                marine_top_winds=sorted([r for r in rankings if r.get('wind_rank')],key=lambda r:r['wind_rank']),
                marine_top_gusts=sorted([r for r in rankings if r.get('gust_rank')],key=lambda r:r['gust_rank']),marine_exclusions=[r for r in rankings if not r['wind_eligible'] or not r['gust_eligible']],
                rainfall_before={'rows':len(previous),'positive':sum(v>0 for v in previous.values())},
                rainfall_after={'rows':len(rain),'positive':sum(r['after']>0 for r in rain)},rainfall_rows=rain)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('output');p.add_argument('baseline');p.add_argument('--report',required=True);a=p.parse_args()
    Path(a.report).write_text(json.dumps(audit(a.output,a.baseline),indent=2,allow_nan=False)+'\n')
