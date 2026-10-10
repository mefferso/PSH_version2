"""Exact CPRA/USACE historical observations from the official MVN CWMS API.

Stage is never qualified by a location's elevation/vertical-datum field or a
requested API datum. Station-specific measured-gauge evidence is mandatory.
"""
import csv
import datetime as dt
import hashlib
import json
import math
import os
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import quote
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from common import UTC, bounds, finite, inventory, public_url, timestamp

BASE='https://cwms-data.usace.army.mil/cwms-data'
MAPPING=Path(__file__).with_name('coastal_water_stations.json')
FILES=['coastal_water_audit.json','coastal_water_review.csv','coastal_water_observations.json']
HEADERS=['Site ID','Network','Current station ID','CWMS series','Peak observed (ft)',
         'Peak UTC','Observed datum / qualification','PSH populated','Flag','Observation count',
         'First UTC','Last UTC','Expected count','Coverage complete','Human review','Reason','Station URL','Historical URL']
FIELDS=['site_id','network','rivergages_sid','cwms_timeseries','peak_ft','peak_time_utc',
        'observed_datum','can_populate_psh','flag','observation_count','first_utc','last_utc',
        'expected_count','coverage_complete','human_review','reason','observation_url','source_url']

def load_mapping():
    rows=json.loads(MAPPING.read_text());mapping={r['site_id']:r for r in rows}
    if len(rows)!=len(mapping):raise ValueError('Duplicate coastal station mapping')
    for m in rows:
        if not m['cwms_timeseries'].startswith(m['cwms_location']+'.Stage.Inst.') or m['cwms_office']!='MVN':
            raise ValueError('Mapping must use exact MVN instantaneous stage series')
        public_url(m['observation_url'])
    return mapping

def match_template(st,wb,m):
    s=wb['Water Level'];r=st['row']
    return (st['id']==m['site_id'] and st['name']==m['template_name'] and st['network']==m['network']
            and s.cell(r,3).value==m['template_latitude'] and s.cell(r,4).value==m['template_longitude']
            and st['url'] in (m['original_url'],m['observation_url']))

def session():
    s=requests.Session()
    retry=Retry(total=2,connect=2,read=2,status=2,backoff_factor=0.5,
                status_forcelist=(429,500,502,503,504),allowed_methods=('GET',),respect_retry_after_header=False)
    s.mount('https://',HTTPAdapter(max_retries=retry))
    s.headers.update({'Accept':'application/json;version=2','User-Agent':'LIX-PSH/2 coastal historical review'})
    return s

def fetch_locations():
    with session() as s:
        r=s.get(BASE+'/locations',params={'office':'MVN'},timeout=(5,25));r.raise_for_status();p=r.json()
    if not isinstance(p,list):raise ValueError('Unexpected CWMS locations schema')
    result={x['name']:x for x in p if x.get('office-id')=='MVN'}
    if len(result)!=len(p):raise ValueError('Duplicate or mismatched CWMS location identities')
    return result

def parse_cwms(p,name,start,end):
    if p.get('name')!=name or p.get('office-id')!='MVN':raise ValueError('CWMS series identity mismatch')
    unit=p.get('units');factor={'ft':1.0,'m':1/0.3048,'cm':1/30.48}.get(unit)
    if factor is None:raise ValueError('CWMS water units not supported explicitly')
    columns={c.get('name'):c.get('ordinal',0)-1 for c in p.get('value-columns',[])}
    if set(columns)!= {'date-time','value','quality-code'} or sorted(columns.values())!=[0,1,2]:
        raise ValueError('CWMS date-time/value/quality schema required')
    if not isinstance(p.get('values'),list):raise ValueError('CWMS values array required')
    rows=[]
    for row in p['values']:
        if not isinstance(row,list) or len(row)!=3:raise ValueError('Malformed CWMS observation')
        ms=finite(row[columns['date-time']],0,253402300799000)
        if ms is None:raise ValueError('Invalid CWMS millisecond timestamp')
        try:t=dt.datetime.fromtimestamp(ms/1000,UTC)
        except (ValueError,OverflowError,OSError) as exc:raise ValueError('Invalid CWMS epoch') from exc
        if not start<=t<end:continue
        raw=row[columns['value']];v=finite(raw);q=row[columns['quality-code']]
        # CWMS 0 = unscreened; 3 = screened/okay. Retain every other code
        # for human review, but never let missing/rejected/estimated values
        # become an automatic measured peak. No interpolation.
        valid=v is not None and q in (0,3) and not isinstance(q,bool) and -100<=v*factor<=100
        reason='' if valid else 'Missing/nonfinite/out-of-range value or non-observed/unqualified quality code'
        rows.append({'time_utc':t.isoformat(),'original_value':v,'original_unit':unit,
                     'value_ft':v*factor if v is not None else None,'quality_code':q,
                     'eligible':valid,'qualification':reason})
    return rows

def summarize(rows,start,end,interval_seconds):
    grouped=defaultdict(list)
    for row in rows:grouped[row['time_utc']].append(row)
    good=[];conflicts=0
    for same in grouped.values():
        values={(r['original_value'],r['original_unit'],r['quality_code']) for r in same}
        if len(values)>1:conflicts+=1;continue
        if same[0]['eligible']:good.append(same[0])
    good.sort(key=lambda r:r['time_utc'])
    peak=max(good,key=lambda r:r['value_ft']) if good else None
    expected=math.ceil((end-start).total_seconds()/interval_seconds)
    # A count alone is insufficient: check every timestamp in the exact grid.
    expected_times={(start+dt.timedelta(seconds=i*interval_seconds)).isoformat() for i in range(expected)}
    actual={r['time_utc'] for r in good}
    return {'peak':peak,'count':len(good),'first_utc':good[0]['time_utc'] if good else None,
            'last_utc':good[-1]['time_utc'] if good else None,'expected_count':expected,
            'missing_count':len(expected_times-actual),'conflicting_times':conflicts,
            'complete':actual==expected_times and not conflicts,
            'quality_codes':sorted({r['quality_code'] for r in rows if isinstance(r['quality_code'],int)})}

def collect(m,start,end,session=None):
    own=session is None;s=session or globals()['session']()
    rows=[];urls=[];errors=[];pages=[];token=None;seen=set();total=None
    params={'office':'MVN','name':m['cwms_timeseries'],'begin':start.isoformat(),'end':end.isoformat(),
            'units':'ft','page-size':500}
    try:
        for _ in range(200):
            if token:params['page']=token
            try:
                r=s.get(BASE+'/timeseries',params=params,timeout=(5,25));r.raise_for_status();p=r.json()
                parsed=parse_cwms(p,m['cwms_timeseries'],start,end)
                if total is None:total=p.get('total',0)
                urls.append(public_url(r.url));rows.extend(parsed)
                pages.append({'url':public_url(r.url),'sha256':hashlib.sha256(json.dumps(p,sort_keys=True).encode()).hexdigest(),'payload':p})
                next_token=p.get('next-page')
                if not next_token:break
                if next_token in seen:raise ValueError('CWMS pagination cycle')
                seen.add(next_token);token=next_token
            except (requests.RequestException,ValueError,TypeError) as exc:
                errors.append(type(exc).__name__+': '+str(exc)[:180]);break
        else:errors.append('CWMS pagination safety limit reached')
    finally:
        if own:s.close()
    coverage=summarize(rows,start,end,m['interval_seconds'])
    if errors:coverage['complete']=False
    return {'observations':rows,'coverage':coverage,'source_urls':urls,'errors':errors,'pages':pages}

def qualify(m,location,start,end):
    rule=m['datum_rule']
    if rule['kind']!='direct':return False,m['notes']
    if (location.get('description') or '')!=m['description'] or location.get('vertical-datum')!=m['location_vertical_datum']:
        return False,'Live gauge metadata changed; datum evidence must be reviewed again'
    effective=timestamp(rule['effective_start_utc'])
    if not effective or start<effective:return False,'Requested observations predate verified direct gauge datum period'
    return True,rule['evidence_text']+'; '+rule['evidence_source']

def collect_hml(m,start,end,session=None):
    """Exact original NWS identifier only; secondary stage is review-only."""
    from datums import HML,parse_hml
    own=session is None;s=session or globals()['session']()
    rows=[];urls=[];errors=[];pages=[]
    try:
        r=s.get(HML,params={'station':m['site_id'],'kind':'obs','tz':'UTC','fmt':'csv',
            'year1':start.year,'month1':start.month,'day1':start.day,
            'year2':end.year,'month2':end.month,'day2':end.day},timeout=(5,25))
        r.raise_for_status()
        pairs=parse_hml(r.text,m['site_id'],start.date(),(end-dt.timedelta(days=1)).date())
        rows=[{'time_utc':t.isoformat(),'original_value':v,'original_unit':'ft','value_ft':v,
               'quality_code':None,'eligible':True,'qualification':'NWS archived stage; datum unverified'} for v,t in pairs]
        urls=[public_url(r.url)]
        pages=[{'url':urls[0],'sha256':hashlib.sha256(r.text.encode()).hexdigest(),'payload_csv':r.text}]
    except (requests.RequestException,ValueError) as exc:errors.append(type(exc).__name__+': HML retrieval/schema unavailable')
    finally:
        if own:s.close()
    coverage=summarize(rows,start,end,m['interval_seconds'])
    # HML lacks source sampling/quality guarantees. Always mark review I.
    coverage['complete']=False
    return {'observations':rows,'coverage':coverage,'source_urls':urls,'errors':errors,'pages':pages,
            'source_kind':'NWS HML archived by IEM (secondary)','datum_eligible':False}

def retrieve(m,start,end):
    result=collect(m,start,end)
    result['source_kind']='USACE MVN CWMS (primary)';result['datum_eligible']=True
    import re
    if not result['coverage']['peak'] and re.fullmatch(r'[A-Z]{4}[0-9]',m['site_id']):
        fallback=collect_hml(m,start,end)
        if fallback['coverage']['peak']:
            fallback['primary_attempt']=result
            return fallback
        result['secondary_attempt']=fallback
    return result

def location_matches(m,loc):
    # Coordinates are checked against the audited source snapshot, with all
    # template discrepancies retained in notes; never pick a nearest gauge.
    return (loc.get('name')==m['cwms_location'] and loc.get('office-id')=='MVN'
            and loc.get('public-name')==m['public_name']
            and loc.get('latitude')==m['source_latitude'] and loc.get('longitude')==m['source_longitude'])

def write_outputs(wb,output_dir,reports,observations):
    output_dir=Path(output_dir);output_dir.mkdir(parents=True,exist_ok=True)
    (output_dir/FILES[0]).write_text(json.dumps(reports,indent=2,allow_nan=False)+'\n')
    (output_dir/FILES[2]).write_text(json.dumps(observations,indent=2,allow_nan=False)+'\n')
    with (output_dir/FILES[1]).open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(HEADERS)
        writer.writerows([[r.get(k) for k in FIELDS] for r in reports])
    if 'Water Level Review' in wb:del wb['Water Level Review']
    sheet=wb.create_sheet('Water Level Review');sheet.append(HEADERS)
    for r in reports:
        sheet.append([r.get(k) for k in FIELDS]);n=sheet.max_row
        for col,key in [(17,'observation_url'),(18,'source_url')]:
            if r.get(key):sheet.cell(n,col).hyperlink=r[key]
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for c in sheet[1]:
        from openpyxl.styles import Font,PatternFill
        c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for col in ('A','B','C','E','F','G','H','I','J','K','L','M','N','O'):sheet.column_dimensions[col].width=24
    for col in ('D','P','Q','R'):sheet.column_dimensions[col].width=65

def populate(wb,qc,start,end,counts,audit,output_dir=Path('output')):
    a,b=bounds(start,end);mapping=load_mapping()
    stations=[s for s in inventory(wb,'Water Level') if s['network'] in ('USACE','LA CPRA')]
    duplicates=Counter(s['id'] for s in stations)
    try:locations=fetch_locations();metadata_error=''
    except (requests.RequestException,ValueError) as exc:locations={};metadata_error='Metadata unavailable: '+type(exc).__name__
    eligible=[s for s in stations if s['id'] in mapping and duplicates[s['id']]==1 and match_template(s,wb,mapping[s['id']])]
    def fetch(st):return st['id'],retrieve(mapping[st['id']],a,b)
    with ThreadPoolExecutor(max_workers=4) as pool:cache=dict(pool.map(fetch,eligible))
    reports=[];observations={}
    registry_path=os.environ.get('PSH_DATUM_REGISTRY')
    entries=json.loads(Path(registry_path).read_text()) if registry_path else []
    if len({e['site_id'] for e in entries})!=len(entries):raise ValueError('Duplicate IDs in datum registry')
    registry={e['site_id']:e for e in entries}
    for st in stations:
        m=mapping.get(st['id']);result=cache.get(st['id']);loc=locations.get(m['cwms_location'],{}) if m else {}
        report={'site_id':st['id'],'row':st['row'],'network':st['network'],'name':st['name'],
                'requested_start_utc':a.isoformat(),'requested_end_utc':b.isoformat(),'human_review':True,
                'can_populate_psh':False,'peak_ft':None,'peak_time_utc':None,'flag':'I',
                'observed_datum':'UNVERIFIED GAGE STAGE','observation_count':0,'coverage_complete':False,
                'source_url':None,'reason':''}
        if m:report.update({k:m[k] for k in ('rivergages_sid','cwms_location','cwms_timeseries','observation_url','original_url','notes','evidence_urls','source_latitude','source_longitude','catalog_extent')})
        if result is None:
            report['reason']='Station mapping/template identity ambiguous or missing';reports.append(report)
            qc.append(['Water Level',st['id'],st['network'],'IDENTITY REVIEW',report['reason'],st['url']]);continue
        observations[st['id']]=dict(result,location_metadata=loc,metadata_url=m['metadata_url'])
        cov=result['coverage'];peak=cov['peak']
        report.update(observation_count=cov['count'],expected_count=cov['expected_count'],first_utc=cov['first_utc'],
                      last_utc=cov['last_utc'],missing_count=cov['missing_count'],coverage_complete=cov['complete'],
                      quality_codes=cov['quality_codes'],conflicting_times=cov['conflicting_times'],errors=result['errors'],
                      source_kind=result.get('source_kind','USACE MVN CWMS (primary)'),
                      source_url=result['source_urls'][0] if result['source_urls'] else m['historical_service'])
        if not peak:
            report['reason']='No eligible observed values in requested window. '+('; '.join(result['errors']) or m['notes'])
            status='NO DATA' if not result['errors'] else 'ERROR'
        else:
            report.update(peak_ft=peak['value_ft'],peak_time_utc=peak['time_utc'],original_value=peak['original_value'],original_unit=peak['original_unit'])
            counts['coastal_water_observed_peaks']+=1
            qualified,reason=qualify(m,loc,a,b)
            if result.get('datum_eligible') is False:qualified=False;reason='Secondary NWS HML stage retained; measured datum not independently established'
            if not location_matches(m,loc):qualified=False;reason=metadata_error or 'Live station identity changed/unverified'
            # Explicit known mismatches never qualify without identity review.
            if st['id'] in ('BBOL1','COCL1'):qualified=False;reason=m['notes']
            value=peak['value_ft'];evidence=m['evidence_urls'];datum='NAVD88' if qualified else 'UNVERIFIED GAGE STAGE'
            if st['id'] in registry:
                from datums import offset
                converted_station=dict(st,url='https://rivergages.mvr.usace.army.mil/WaterControl/stationinfo2.cfm?sid='+m['rivergages_sid'])
                adjustment=offset(registry[st['id']],converted_station,a,b)
                if not location_matches(m,loc) or st['id'] in ('BBOL1','COCL1'):
                    raise ValueError('Datum registry cannot resolve station identity conflict')
                value+=adjustment;qualified=True;datum='NAVD88';evidence=registry[st['id']]['evidence']
                reason=f'Independently human-reviewed event-effective stage + {adjustment} ft conversion'
            report['observed_datum']=datum if qualified else ('NGVD29 stage' if st['id'] in ('BBOL1','TSPL1','BDAL1','BCSL1') else 'UNVERIFIED GAGE STAGE')
            coverage_note=f'{cov["count"]}/{cov["expected_count"]} expected observations; first {cov["first_utc"]}, last {cov["last_utc"]}; missing {cov["missing_count"]}; conflicts {cov["conflicting_times"]}'
            report['reason']=reason+'; '+coverage_note+'; '+m['notes']
            if qualified:
                if finite(value,-100,100) is None:raise ValueError('Converted elevation outside QC range')
                s=wb['Water Level'];r=st['row'];t=timestamp(peak['time_utc'])
                s.cell(r,7).value=round(value,2);s.cell(r,8).value=datum
                for n,v in enumerate((t.strftime('%H%M'),t.day,t.month,t.year)):s.cell(r,9+n).value=v
                s.cell(r,14).value='I' if not cov['complete'] else None
                audit.add('Water Level',r,st['id'],'water',round(value,2),'ft',t,report['source_url'],datum=datum,
                          evidence=evidence,raw_value=peak['original_value'],raw_unit=peak['original_unit'],details=report['reason'])
                report.update(can_populate_psh=True,psh_value_ft=round(value,2),flag='I' if not cov['complete'] else '')
                counts['coastal_water_psh_peaks']+=1;status='INCOMPLETE' if not cov['complete'] else 'REVIEW REQUIRED'
            else:status='DATUM REVIEW';counts['coastal_water_datum_review']+=1
        qc.append(['Water Level',st['id'],st['network'],status,report['reason'],report['source_url']]);reports.append(report)
    write_outputs(wb,output_dir,reports,observations)
    return reports
