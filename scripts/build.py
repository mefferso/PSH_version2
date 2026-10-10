"""PSH review pipeline: exact template inventory, safe collection and audited products."""
import datetime as dt
import json
import os
from pathlib import Path
import re
from collections import Counter
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font,PatternFill
from common import Audit,inventory,bounds
from cocorahs import rain_bounds
import asos1min,iem,hourlyprecip,coopobs,coops,ndbc,usgs,usace,synoptic,weatherstem,cocorahs,hads,tornadoes,imports,products,datums

TEMPLATE=Path('Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')
OUT=Path('output')

def collectors():
    return [hourlyprecip.populate,asos1min.populate,iem.populate,coops.populate,ndbc.populate,usgs.populate,synoptic.populate,
            weatherstem.populate,cocorahs.populate,synoptic.populate_rain,hads.populate,coopobs.populate,tornadoes.populate,datums.populate]

def reset(wb):
    for tab,cols in [('Wind and Pressure',range(11,31)),('Rainfall',range(8,10)),('Water Level',[7,9,10,11,12,14])]:
        s=wb[tab]
        for st in inventory(wb,tab):
            if st['network'].upper() in ('WEATHERFLOW','WXFLOW'):continue
            for col in cols:s.cell(st['row'],col).value=None
    s=wb['Tornadoes']
    for r in range(2,s.max_row+1):
        if str(s.cell(r,1).value or '').startswith('[Insert'):
            for c in range(1,12):
                if not isinstance(s.cell(r,c),MergedCell):s.cell(r,c).value=None
    # Clear unverified example counts/narratives; preserve definitions and county roster.
    for c in ('B9','B11','A14'):wb['Summary'][c]=None
    flood=wb['Inland Flooding'];flood['B3']=None
    for r in range(6,flood.max_row+1):
        if not isinstance(flood.cell(r,4),MergedCell):flood.cell(r,4).value=None
    s=wb['Impacts']
    for r in range(6,s.max_row+1):
        for c in range(3,s.max_column+1):
            if not isinstance(s.cell(r,c),MergedCell):s.cell(r,c).value=None

def coverage(wb,qc,counts,audit):
    recorded={(r[0].value,str(r[1].value)) for r in list(qc)[1:]}
    imported={(x['tab'],x['site_id']) for x in audit.entries if x['status']=='REVIEWED IMPORT'}
    for tab in ('Wind and Pressure','Rainfall','Water Level'):
        for st in inventory(wb,tab):
            key=(tab,st['id'])
            if key in imported:
                qc.append([tab,st['id'],st['network'],'REVIEWED IMPORT','Explicitly reviewed import; see provenance',st['url']]);continue
            if key in recorded:continue
            manual=st['network'].upper() in ('WEATHERFLOW','WXFLOW')
            detail='Automatic WeatherFlow collection intentionally excluded; metadata and manual values retained' if manual else {
                'USGS':'No validated rain/wind series for this inventory row; direct water collector is separate',
                'USACE':'Historical series and event-effective datum require verification',
                'LA CPRA':'Historical series and event-effective datum require verification',
                'TPCG':'No source hyperlink/archive station mapping in inventory',
                'COOP':'Daily rain reporting periods and exact archive mapping require validation; no calendar-day substitution',
                'HADS':'Precipitation counter/increment semantics and interval coverage require validation',
                'RAWS':'Rainfall accumulation series and interval coverage require validation',
            }.get(st['network'],'No validated collector for this exact inventory row')
            status='MANUAL' if manual else 'UNSUPPORTED'
            qc.append([tab,st['id'],st['network'],status,detail,st['url']]);counts[status]+=1
    for tab in ('Inland Flooding','Impacts'):
        qc.append([tab,'','','MANUAL REVIEW','No verified storm-specific impact narrative; unknown fatalities/injuries/evacuations remain blank',''])

def build():
    name=os.environ['STORM_NAME'].strip();start=dt.date.fromisoformat(os.environ['START_UTC']);end=dt.date.fromisoformat(os.environ['END_UTC'])
    if not re.fullmatch(r'[A-Za-z0-9 -]{2,70}',name):raise ValueError('Invalid storm name')
    if start>end or (end-start).days>35:raise ValueError('Invalid date range (maximum 36 inclusive days)')
    wb=load_workbook(TEMPLATE)
    required={'Summary','Wind and Pressure','Rainfall','Water Level','Tornadoes','Impacts','Inland Flooding'}
    if not required.issubset(wb.sheetnames):raise ValueError('Missing PSH template tabs')
    reset(wb)
    if 'QC' in wb:del wb['QC']
    qc=wb.create_sheet('QC');qc.append(['Tab','Site ID','Network','Status','Details','Source URL'])
    counts=Counter();audit=Audit()
    wb['Summary']['B3']=name;wb['Summary']['B5']='NWS New Orleans/Baton Rouge'
    wb['Summary']['B7']=f'{start:%m/%d/%Y} - {end:%m/%d/%Y}'
    updated=dt.datetime.now(dt.timezone.utc).strftime('%m/%d/%Y')
    for tab,row,detailrow in [('Wind and Pressure',172,173),('Rainfall',230,231),('Water Level',115,116)]:
        wb[tab].cell(row,2).value=updated
        wb[tab].cell(detailrow,2).value='Automated review build; source availability and qualifications are listed in QC and provenance.'
    wb['Summary']['A113']='Review workbook last generated on '+updated
    wb['Wind and Pressure']['B174']='Source sampling, averaging periods, exposure and coverage require review; qualified source values only. Unavailable values remain blank.'
    wb['Water Level']['B117']='NOS readings use verified MHHW metadata. Other water values require direct NAVD88 or independently reviewed event-effective conversion evidence. NAVD88 elevation is not automatically inundation depth. Sampling and source availability vary; see QC.'
    a,b=rain_bounds(start,end)
    wb['Rainfall']['B232']=a.strftime('%H%M UTC %b %d %Y');wb['Rainfall']['B233']=b.strftime('%H%M UTC %b %d %Y')
    wb['Rainfall']['B234']='Only fully documented accumulation windows are automated. No interpolation across missing periods or inclusion of overlapping daily reports.'
    wb['Summary']['A117']='Development review product. Storm label does not establish cyclone attribution. Station exposure, observation coverage and datums require review. Candidate CSV thresholds follow NWSI 10-601 (2026-08-17).'
    wb['Summary']['A114']='Automated review outputs; see QC and provenance. No official PSH issuance.'
    # Always manifest the partial-observation CSV, including offline fixture
    # builds whose mocked collector list intentionally omits CoCoRaHS.
    import csv
    OUT.mkdir(parents=True,exist_ok=True)
    with (OUT/'Rainfall_partial_reports.csv').open('w',newline='',encoding='utf-8') as fh:
        csv.writer(fh).writerow(['station_id','report_type','period_start_utc',
            'period_end_utc','reported_in','requested_start_utc','requested_end_utc',
            'classification','source_url'])
    with (OUT/'USGS_stage_review.csv').open('w',newline='',encoding='utf-8') as fh:
        csv.writer(fh).writerow(['site_id','usgs_site_number','peak_gage_height_ft',
            'peak_time_utc','datum','qualification','source_url'])
    for adapter in collectors():
        if adapter in (cocorahs.populate,hads.populate,coopobs.populate,usgs.populate,iem.populate):
            adapter(wb,qc,start,end,counts,audit,output_dir=OUT)
        else:
            adapter(wb,qc,start,end,counts,audit)
    if os.environ.get('PSH_OFFLINE')!='1':usace.populate(wb,qc,start,end,counts)
    if os.environ.get('PSH_IMPORT_FILE'):imports.load(wb,os.environ['PSH_IMPORT_FILE'],start,end,audit)
    coverage(wb,qc,counts,audit)
    products.summaries(wb)
    OUT.mkdir(parents=True,exist_ok=True)
    issues=products.export_csv(wb,OUT/'csv')
    for issue in issues:qc.append(['CSV','','','DUPLICATE ID',issue,''])
    qc.freeze_panes='A2';qc.auto_filter.ref=f'A1:F{qc.max_row}'
    for c in qc[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for col,width in zip('ABCDEF',[22,20,20,24,100,70]):qc.column_dimensions[col].width=width
    target=OUT/f'PSHLIX_{start.year}_{re.sub(r"[^A-Za-z0-9_-]+","_",name)}_REVIEW.xlsx';wb.save(target)
    imports.template(wb,OUT/'WeatherFlow-import-template.json');audit.save(OUT/'provenance.json')
    import hashlib
    statuses=Counter(str(qc.cell(r,4).value) for r in range(2,qc.max_row+1))
    report={'storm':name,'start_utc':str(start),'end_utc':str(end),'rain_start_utc':a.isoformat(),'rain_end_utc':b.isoformat(),
            'generated_at_utc':dt.datetime.now(dt.timezone.utc).isoformat(),'coverage':'DEVELOPMENT REVIEW — incomplete source coverage',
            'csv_files':products.CSV_FILES,'workbook':target.name,'measurement_count':len(audit.entries),'counts':dict(counts),'qc_status_counts':dict(statuses),
            'manual_networks':['WeatherFlow'],'csv_kind':'REVIEW and CANDIDATE — NWSI 10-601 (2026-08-17) thresholds; not official issuance',
            'template_sha256':hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),'commit':os.environ.get('GITHUB_SHA','local'),
            'workflow_url':('https://github.com/'+os.environ['GITHUB_REPOSITORY']+'/actions/runs/'+os.environ['GITHUB_RUN_ID']) if os.environ.get('GITHUB_RUN_ID') else None,
            'validation':'Live accuracy and Francine source regression required; parser/integration tests alone are insufficient',
            'inventory_counts':{tab:len(list(inventory(wb,tab))) for tab in ('Wind and Pressure','Rainfall','Water Level')}}
    (OUT/'QC.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2));return report

if __name__=='__main__':build()
