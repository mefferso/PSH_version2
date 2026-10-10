"""Native template summaries and explicitly nonofficial review CSV exports."""
import csv
from collections import Counter
from pathlib import Path
from openpyxl.cell.cell import MergedCell
from common import inventory,finite

CSV_FILES=['Wind_and_Pressure_REVIEW.csv','Rainfall_REVIEW.csv','Water_Level_REVIEW.csv','WindandPressure_CANDIDATE.csv','Rainfall_CANDIDATE.csv','WaterLevel_CANDIDATE.csv','Tornadoes_CANDIDATE.csv']

BLOCKS=[(20,'Wind and Pressure',11,True,'L'),(34,'Wind and Pressure',17,True,'L'),
        (48,'Wind and Pressure',11,True,'M'),(62,'Wind and Pressure',17,True,'M'),
        (76,'Rainfall',8,True,None),(89,'Water Level',7,True,'NOS'),
        (102,'Wind and Pressure',23,False,None)]

def csv_values(values):
    return [int(v) if isinstance(v,float) and v.is_integer() else v for v in values]

def summaries(wb):
    target=wb['Summary']
    from observation_review import marine
    marine_rows={r['row']:r for r in marine(wb)}
    for first,tab,col,descending,kind in BLOCKS:
        for row in target.iter_rows(min_row=first,max_row=first+9,min_col=1,max_col=4):
            for cell in row:
                if not isinstance(cell,MergedCell):cell.value=None;cell.hyperlink=None
        s=wb[tab];eligible=[];seen=set()
        for station in inventory(wb,tab):
            r=station['row'];value=finite(s.cell(r,col).value)
            if value is None:continue
            if kind=='M':
                record=marine_rows.get(r)
                if not record or not record[('wind' if col==11 else 'gust')+'_eligible']:continue
            elif kind=='L':
                height=finite(s.cell(r,9).value,0,1000)
                if str(s.cell(r,7).value or '')!=kind or height is None or height>=20:continue
            if kind=='NOS' and station['network']!='NOS':continue
            # Same exact observing source only once in a top-10 block.
            key=(station['id'],station['url'])
            if key in seen:continue
            seen.add(key);eligible.append((value,station))
        eligible.sort(key=lambda x:(-x[0] if descending else x[0],x[1]['id'],x[1]['row']))
        for offset,(value,station) in enumerate(eligible[:10]):
            r=station['row']
            if kind=='M':values=[station['name'],station['network'],value]
            elif tab=='Water Level':values=[station['name'],s.cell(r,6).value,s.cell(r,8).value,value]
            else:values=[station['name'],s.cell(r,6).value,station['network'],value]
            for c,v in enumerate(values,1):target.cell(first+offset,c).value=v
            if station['url']:target.cell(first+offset,1).hyperlink=station['url']
    # Ensure any residual Google dummy formulas cannot show cached historical examples.
    for row in target:
        for c in row:
            if c.data_type=='f' and '__xludf' in c.value:c.value=None

def export_csv(wb,directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True);issues=[]
    for tab,cols,measurements in [('Wind and Pressure',30,(11,17,23)),('Rainfall',9,(8,)),('Water Level',14,(7,))]:
        s=wb[tab];stations=list(inventory(wb,tab));duplicates=Counter(x['id'] for x in stations)
        target=directory/(tab.replace(' ','_')+'_REVIEW.csv')
        with target.open('w',newline='',encoding='utf-8') as stream:
            writer=csv.writer(stream);writer.writerow([s.cell(1,c).value or f'Column {c}' for c in range(1,cols+1)])
            for station in stations:
                r=station['row']
                if duplicates[station['id']]>1:
                    issues.append(f'{tab}: duplicate ID {station["id"]} at row {r}; excluded from review CSV')
                    continue
                if not any(finite(s.cell(r,c).value) is not None for c in measurements):continue
                values=[s.cell(r,c).value for c in range(1,cols+1)];values[0]=station['id']
                writer.writerow(csv_values(values))
    # NWSI 10-601 (2026-08-17), section 8: gust >33 kn, MSLP <1005 mb,
    # rain >=3 inches. These candidate files still require meteorologist review.
    for tab,filename,cols in (("Wind and Pressure","WindandPressure",29),("Rainfall","Rainfall",9),("Water Level","WaterLevel",14)):
        source=wb[tab];stations=list(inventory(wb,tab));ids=Counter(st["id"] for st in stations)
        with (directory/(filename+"_CANDIDATE.csv")).open("w",newline="",encoding="utf-8") as stream:
            writer=csv.writer(stream);writer.writerow([source.cell(1,c).value for c in range(1,cols+1)])
            for st in stations:
                if ids[st["id"]]>1 or not st["network"]:continue
                r=st["row"]
                gust=finite(source.cell(r,17).value) if tab=="Wind and Pressure" else None
                pressure=finite(source.cell(r,23).value) if tab=="Wind and Pressure" else None
                rain=finite(source.cell(r,8).value) if tab=="Rainfall" else None
                water=finite(source.cell(r,7).value) if tab=="Water Level" else None
                eligible=((gust is not None and gust>33) or (pressure is not None and pressure<1005)) if tab=="Wind and Pressure" else rain is not None and rain>=3 if tab=="Rainfall" else water is not None
                if not eligible:continue
                values=[source.cell(r,c).value for c in range(1,cols+1)];values[0]=st["id"];writer.writerow(csv_values(values))
    tornado=wb["Tornadoes"]
    with (directory/"Tornadoes_CANDIDATE.csv").open("w",newline="",encoding="utf-8") as stream:
        writer=csv.writer(stream);writer.writerow([tornado.cell(1,c).value for c in range(1,12)])
        for r in range(2,tornado.max_row+1):
            if finite(tornado.cell(r,2).value,-90,90) is not None and finite(tornado.cell(r,3).value,-180,180) is not None:
                writer.writerow(csv_values([tornado.cell(r,c).value for c in range(1,12)]))
    return issues
