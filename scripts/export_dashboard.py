"""Manifest-selected workbook, all PSH tabs, measurement audit and review downloads."""
import datetime as dt
import json
from pathlib import Path
import shutil
import zipfile
from openpyxl import load_workbook

def clean(v):return v.isoformat() if isinstance(v,(dt.datetime,dt.date)) else v

def export(out=Path('output'),site=Path('site')):
    out=Path(out);site=Path(site);meta=json.loads((out/'QC.json').read_text())
    target=out/meta['workbook'];book=load_workbook(target)
    audit=json.loads((out/'provenance.json').read_text());result={'meta':meta,'tabs':{},'download':target.name,'audit':audit}
    for s in book:
        header_row=5 if s.title in ('Impacts','Inland Flooding') else 1
        headers=[clean(s.cell(header_row,c).value) for c in range(1,s.max_column+1)]
        rows=[]
        for r in range(1 if s.title=='Summary' else header_row+1,s.max_row+1):
            values=[clean(s.cell(r,c).value) for c in range(1,s.max_column+1)]
            if not any(v is not None for v in values):continue
            if s.title!='Summary' and (values[0] is None or str(values[0]).startswith('[Insert')):continue
            links={str(c-1):s.cell(r,c).hyperlink.target for c in range(1,s.max_column+1) if s.cell(r,c).hyperlink and str(s.cell(r,c).hyperlink.target).startswith(('https://','http://'))}
            rows.append({'v':values,'links':links,'row':r})
        result['tabs'][s.title]={'status':'REVIEW','headers':headers,'rows':rows}
    result['tabs']['Provenance']={'status':'REVIEW','headers':['Tab','Row','Site ID','Variable','Value','Unit','UTC time','Datum','Status','Source URL','Details'],
        'rows':[{'v':[e.get(k) for k in ('tab','row','site_id','variable','value','unit','time_utc','datum','status','source_url','details')],
                 'links':{'9':e['source_url']}} for e in audit]}
    from common import identifier
    from coastal_water import FILES
    review=json.loads((out/FILES[0]).read_text());result['water_review']=review
    by_id={(r['site_id'],r['row']):r for r in review}
    water=result['tabs']['Water Level'];base=len(water['headers'])
    water['headers'].extend(['Review peak (ft; original datum)','Review peak UTC','Review datum','Coverage / qualification','Historical observations'])
    water['notice']='PSH water values require verified datum evidence. Review peaks retain measured stage in its original datum; I identifies incomplete or unqualified observations. Human meteorologist review required.'
    for row in water['rows']:
        r=by_id.get((identifier(row['v'][0]),row['row']));row['v']+=[None]*5
        if r:
            row['v'][base:]=[r.get('peak_ft'),r.get('peak_time_utc'),r.get('observed_datum'),r.get('reason'),r.get('source_url')]
            if r.get('source_url'):row['links'][str(base+4)]=r['source_url']
    from observation_review import FILES as REVIEW_FILES
    # Optional only for legacy fixture producers; production build always writes these.
    for name in REVIEW_FILES:
        if (out/name).exists():
            if name.endswith('.json'):result[name.removesuffix('.json')]=json.loads((out/name).read_text())
    for tab in ('Summary','Wind and Pressure','Rainfall','Water Level'):
        result['tabs'][tab]['notice']=result['tabs'][tab].get('notice','')+' Meteorologist review required. '+('PROVISIONAL: requested event period has not ended. ' if meta.get('provisional') else '')
    wind=result['tabs']['Wind and Pressure'];base_wind=len(wind['headers'])
    wind['headers'].extend(['Marine sustained Summary eligibility','Marine gust Summary eligibility'])
    marine_index={r['row']:r for r in result.get('marine_summary_review',[])}
    for row in wind['rows']:
        rec=marine_index.get(row['row']);row['v']+=[rec.get('wind_reason') if rec else None,rec.get('gust_reason') if rec else None]
    rain=result['tabs']['Rainfall'];rain['headers'].append('Accumulation review')
    rain_index={r['row']:r for r in result.get('observation_coverage_review',[]) if r['tab']=='Rainfall'}
    for row in rain['rows']:
        rec=rain_index.get(row['row']);row['v'].append((rec['interpretation']+'; '+rec['qualification']) if rec else None)
    result['tabs']['Marine Summary Review']={'status':'REVIEW','headers':['ID','Name','Network','Height m','Mean min','Sustained kn','Wind rank','Wind reason','Gust kn','Gust rank','Gust reason'],
        'rows':[{'v':[r.get(k) for k in ('id','name','network','height_m','averaging_minutes','sustained_kn','wind_rank','wind_reason','gust_kn','gust_rank','gust_reason')],'links':{'0':r['url']} if r['url'] else {}} for r in result.get('marine_summary_review',[])]}
    result['tabs']['Rainfall Archive Review']={'status':'REVIEW','notice':'Individual source reports with unknown UTC observation clocks/periods; not storm totals. Primary source download failures remain in QC.','headers':['ID','Name','Source date (not UTC observation time)','Reported in','Qualification','Source URL'],'rows':[{'v':[r.get(k) for k in ('site_id','template_name','source_date','reported_inches','qualification','source_url')],'links':{'5':r['source_url']}} for r in result.get('rainfall_archive_review',{}).get('reports',[])]}
    site.mkdir(parents=True,exist_ok=True);(site/'data').mkdir(exist_ok=True)
    (site/'data/latest.json').write_text(json.dumps(result,ensure_ascii=False,allow_nan=False))
    shutil.copyfile(target,site/target.name)
    with zipfile.ZipFile(site/'review-outputs.zip','w',zipfile.ZIP_DEFLATED) as z:
        paths=[target,out/'QC.json',out/'provenance.json',out/'WeatherFlow-import-template.json',out/'Rainfall_partial_reports.csv',out/'USGS_stage_review.csv']
        paths.extend(out/name for name in FILES)
        paths.extend(out/name for name in REVIEW_FILES if (out/name).exists())
        paths.extend(out/'csv'/name for name in meta.get('csv_files',[]))
        for path in paths:
            if not path.is_file():raise ValueError('Missing manifested artifact: '+str(path))
            z.write(path,path.relative_to(out))
    shutil.copyfile(out/'WeatherFlow-import-template.json',site/'WeatherFlow-import-template.json')
    shutil.copyfile(out/'Rainfall_partial_reports.csv',site/'Rainfall_partial_reports.csv')
    shutil.copyfile(out/'USGS_stage_review.csv',site/'USGS_stage_review.csv')
    for name in REVIEW_FILES:
        if (out/name).exists():shutil.copyfile(out/name,site/name)
    for name in FILES:shutil.copyfile(out/name,site/name)
    (site/'.nojekyll').write_text('')
    print(f'Dashboard: {len(book.sheetnames)} workbook tabs + provenance; {len(audit)} measurements');return result

if __name__=='__main__':export()
