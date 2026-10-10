"""Fail publication on changed station metadata or unaudited/mistimed measurements."""
import datetime as dt
from copy import copy,deepcopy
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from common import inventory,identifier,finite,timestamp,bounds

COLS={'Wind and Pressure':{'wind':11,'gust':17,'pressure':23},'Rainfall':{'rain':8},'Water Level':{'water':7}}

def validate(out=Path('output'),template=Path('Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx')):
    out=Path(out);meta=json.loads((out/'QC.json').read_text());entries=json.loads((out/'provenance.json').read_text())
    name=meta['workbook']
    if Path(name).name!=name:raise ValueError('Unsafe workbook path')
    if meta.get('template_sha256')!=hashlib.sha256(Path(template).read_bytes()).hexdigest():raise ValueError('Template checksum mismatch')
    original=load_workbook(template);w=load_workbook(out/name)
    if w.sheetnames!=original.sheetnames+['QC','Water Level Review']:raise ValueError('Original tabs changed')
    for tab,n in [('Wind and Pressure',10),('Rainfall',7),('Water Level',6)]:
        old=original[tab];s=w[tab]
        if set(str(x) for x in old.merged_cells.ranges)!=set(str(x) for x in s.merged_cells.ranges):raise ValueError('Merged template structure changed')
        for row in [1]+[st["row"] for st in inventory(original,tab)]:
            for col in range(1,n+1):
                x=old.cell(row,col);y=s.cell(row,col)
                if x.value!=y.value or any(copy(getattr(x,k))!=copy(getattr(y,k)) for k in ("font","fill","border","alignment","protection","number_format")) or (x.hyperlink.target if x.hyperlink else None)!=(y.hyperlink.target if y.hyperlink else None):
                    raise ValueError(f'Station metadata/style/link changed: {tab} {x.coordinate}')
    if len(entries)!=meta['measurement_count']:raise ValueError('Audit count mismatch')
    audit_map={}
    start,end=bounds(dt.date.fromisoformat(meta['start_utc']),dt.date.fromisoformat(meta['end_utc']))
    rain_start=timestamp(meta.get('rain_start_utc'));rain_end=timestamp(meta.get('rain_end_utc'))
    if not rain_start or not rain_end or rain_start>=rain_end:raise ValueError('Invalid manifest rainfall window')
    for e in entries:
        key=(e['tab'],e['row'],e['variable'])
        if key in audit_map:raise ValueError(f'Duplicate measurement audit: {key}')
        audit_map[key]=e
        if finite(e['value']) is None or not timestamp(e['time_utc']):raise ValueError('Invalid audit value/time')
        if e['variable']=='rain':
            if timestamp(e.get('interval_start_utc'))!=rain_start or timestamp(e['time_utc'])!=rain_end:raise ValueError('Rain audit interval differs from manifest')
        elif not start<=timestamp(e['time_utc'])<end:raise ValueError('Audit time outside window')
        if e['variable']=='water' and (e['datum'] not in ('NAVD88','MHHW') or not e['datum_evidence']):raise ValueError('Unverified water datum')
        if any(x in e['source_url'].lower() for x in ('token=','api_key=','apikey=')):raise ValueError('Credential in audit URL')
    for tab,columns in COLS.items():
        s=w[tab]
        for st in inventory(w,tab):
            for variable,col in columns.items():
                value=s.cell(st['row'],col).value
                if value is None:continue
                if finite(value) is None:raise ValueError('Nonnumeric measurement')
                e=audit_map.get((tab,st['row'],variable))
                if not e:
                    # Only unchanged original WeatherFlow manual readings may lack audit.
                    if st['network'].upper()=='WEATHERFLOW' and value==original[tab].cell(st['row'],col).value:continue
                    raise ValueError(f'No provenance: {tab} {st["id"]} {variable}')
                if e['value']!=value or e['site_id']!=st['id']:raise ValueError('Audit/workbook mismatch')
                if variable=='water' and e['datum']!=s.cell(st['row'],8).value:raise ValueError('Audit datum mismatch')
                if variable!='rain':
                    t=timestamp(e['time_utc']);offset=col+2 if variable in ('wind','gust') else col+1 if variable=='pressure' else 9
                    actual=[s.cell(st['row'],offset+n).value for n in range(4)]
                    if actual!=[t.strftime('%H%M'),t.day,t.month,t.year]:raise ValueError('UTC timestamp/workbook mismatch')
    # Reject orphan audit entries and validate confirmed tornado timestamps/locations.
    for e in entries:
        if e['tab']=='Tornadoes':
            r=e['row'];t=timestamp(e['time_utc']);sheet=w['Tornadoes']
            if e['variable']!='tornado' or e['value']!=1 or finite(sheet.cell(r,2).value,-90,90) is None or finite(sheet.cell(r,3).value,-180,180) is None:raise ValueError('Invalid tornado audit/row')
            if [sheet.cell(r,c).value for c in range(6,10)]!=[t.strftime('%H%M'),t.day,t.month,t.year]:raise ValueError('Tornado UTC mismatch')
        else:
            col=COLS.get(e['tab'],{}).get(e['variable'])
            if not col or w[e['tab']].cell(e['row'],col).value!=e['value']:raise ValueError('Orphan measurement audit')
    from coastal_water import FILES,HEADERS,FIELDS
    review=json.loads((out/FILES[0]).read_text())
    if len({r['site_id'] for r in review})!=len(review):raise ValueError('Duplicate water review ID')
    review_sheet=w['Water Level Review']
    if [c.value for c in review_sheet[1]]!=HEADERS:raise ValueError('Water review headers mismatch')
    for n,r in enumerate(review,2):
        actual=[review_sheet.cell(n,c+1).value for c in range(len(FIELDS))]
        expected=[r.get(k) if r.get(k)!='' else None for k in FIELDS]
        for x,y in zip(actual,expected):
            if isinstance(x,(int,float)) and isinstance(y,(int,float)):
                if abs(x-y)>1e-10:raise ValueError('Water review workbook mismatch')
            elif x!=y:raise ValueError('Water review workbook mismatch')
        if r.get('peak_ft') is not None and (finite(r['peak_ft']) is None or not timestamp(r.get('peak_time_utc')) or not start<=timestamp(r['peak_time_utc'])<end):raise ValueError('Invalid water review peak/time')
        if r.get('can_populate_psh'):
            e=audit_map.get(('Water Level',r['row'],'water'))
            if not e or e['site_id']!=r['site_id'] or e['time_utc']!=r['peak_time_utc'] or e['value']!=r['psh_value_ft']:raise ValueError('Qualified review lacks matching water provenance')
    import products
    expected=deepcopy(w);products.summaries(expected)
    for first,*_ in products.BLOCKS:
        for r in range(first,first+10):
            for c in range(1,5):
                x=w['Summary'].cell(r,c);y=expected['Summary'].cell(r,c)
                if x.value!=y.value or (x.hyperlink.target if x.hyperlink else None)!=(y.hyperlink.target if y.hyperlink else None):raise ValueError('Summary differs from station measurements')
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        products.export_csv(w,Path(d))
        if meta.get('csv_files')!=products.CSV_FILES:raise ValueError('Unexpected CSV artifact list')
        for name in products.CSV_FILES:
            if (out/'csv'/name).read_bytes()!=(Path(d)/name).read_bytes():raise ValueError('CSV differs from validated workbook: '+name)
    print(f'Validated {len(entries)} auditable measurements and preserved station metadata');return True

def validate_site(out=Path('output'),site=Path('site')):
    import zipfile
    out=Path(out);site=Path(site);meta=json.loads((out/'QC.json').read_text())
    payload=json.loads((site/'data/latest.json').read_text())
    if payload['meta']!=meta or payload['audit']!=json.loads((out/'provenance.json').read_text()):raise ValueError('Dashboard manifest/audit mismatch')
    # Re-export independently to compare all tab values, hyperlinks and archive members.
    import tempfile
    from export_dashboard import export
    with tempfile.TemporaryDirectory() as d:
        expected=export(out,Path(d))
        if payload!=expected:raise ValueError('Dashboard workbook values differ')
    files=[meta['workbook'],'QC.json','provenance.json','WeatherFlow-import-template.json','Rainfall_partial_reports.csv','USGS_stage_review.csv']+['csv/'+n for n in meta['csv_files']]
    from coastal_water import FILES
    files.extend(FILES)
    from observation_review import FILES as REVIEW_FILES
    files.extend(name for name in REVIEW_FILES if (out/name).exists())
    with zipfile.ZipFile(site/'review-outputs.zip') as z:
        if sorted(z.namelist())!=sorted(files):raise ValueError('Archive contains unrelated/missing files')
        for name in files:
            if z.read(name)!=(out/name).read_bytes():raise ValueError('Archive content mismatch')
    if (site/meta['workbook']).read_bytes()!=(out/meta['workbook']).read_bytes():raise ValueError('Workbook download mismatch')
    if (site/'Rainfall_partial_reports.csv').read_bytes()!=(out/'Rainfall_partial_reports.csv').read_bytes():raise ValueError('CoCoRaHS partial download mismatch')
    if (site/'USGS_stage_review.csv').read_bytes()!=(out/'USGS_stage_review.csv').read_bytes():raise ValueError('USGS stage review download mismatch')
    for name in FILES:
        if (site/name).read_bytes()!=(out/name).read_bytes():raise ValueError('Coastal water download mismatch')
    print('Validated dashboard tabs and storm-specific downloads')

if __name__=='__main__':
    import sys
    validate()
    if '--site' in sys.argv:validate_site()
