"""NDBC standard meteorological observations; headers determine units/columns."""
import datetime as dt
import gzip
import math
import re
import time
from urllib.parse import urlparse, parse_qs
import requests
from openpyxl.cell.cell import MergedCell
from common import identifier, inventory, observation_now

KNOTS_PER_MS = 1.9438444924406
ALLOWED = {"BUOY", "CMAN", "C-MAN", "WLON", "NDBC"}
BASE = "https://www.ndbc.noaa.gov"
HEADERS = {"YY","MM","DD","hh","mm","WDIR","WSPD","GST","PRES"}

def station_id(cell):
    link = cell.hyperlink.target if cell.hyperlink else ""
    url = urlparse(link or "")
    if url.hostname not in ("ndbc.noaa.gov", "www.ndbc.noaa.gov"):
        return None
    match = re.search(r"/(?:realtime2|stdmet)/([A-Za-z0-9]{4,7})\.(?:txt|gz)$", url.path, re.I)
    value = match.group(1) if match else parse_qs(url.query).get("station", [""])[0]
    if not value:
        value = parse_qs(url.query).get("station_id", [""])[0]
    return value.upper() if re.fullmatch(r"[A-Za-z0-9]{4,7}", value or "") else None

def parse_text(content, start, end):
    headings = None
    units_verified = False
    obs = []
    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue
        fields = line.lstrip("#").split()
        normalized = [x.lstrip("#") for x in fields]
        if "WSPD" in normalized and "PRES" in normalized and ("YY" in normalized or "YYYY" in normalized):
            headings = normalized
            continue
        if headings is not None and line.startswith("#") and "m/s" in fields:
            units = dict(zip(headings,fields))
            if units.get("WSPD") != "m/s" or units.get("GST") != "m/s" or units.get("PRES") not in ("hPa","mb"):
                raise ValueError("Unsupported NDBC units")
            units_verified = True
            continue
        if headings is None or line.startswith("#"):
            continue
        if not units_verified:
            raise ValueError("NDBC unit row missing or unsupported")
        if len(fields) < len(headings):
            continue
        row = dict(zip(headings, fields))
        try:
            yr = int(row.get("YYYY", row.get("YY")))
            if yr < 100: yr += 2000 if yr < 70 else 1900
            minute = int(row.get("mm", "0"))
            moment = dt.datetime(yr, int(row["MM"]), int(row["DD"]), int(row["hh"]), minute,
                                 tzinfo=dt.timezone.utc)
        except (ValueError, KeyError, TypeError):
            continue
        if not (start <= moment.date() <= end) or moment>=observation_now():
            continue
        def number(key, lo, hi):
            try:
                v = float(row[key])
                return v if math.isfinite(v) and lo <= v <= hi else None
            except (KeyError, ValueError, TypeError):
                return None
        # NDBC WSPD/GST m/s and PRES hPa; do not treat missing 99/999 sentinels as data.
        wind = number("WSPD", 0, 75)
        gust = number("GST", 0, 90)
        pressure = number("PRES", 850, 1100)
        direction = number("WDIR", 0, 360)
        obs.append({"time":moment, "wind":wind*KNOTS_PER_MS if wind is not None else None,
                    "gust":gust*KNOTS_PER_MS if gust is not None else None,
                    "pressure":pressure, "dir":direction})
    if headings is None:
        raise ValueError("NDBC standard-meteorological column header not found")
    return obs

def retrieve(station, start, end, session=requests):
    # Current/recent 45 days: realtime2; older and future archive windows: annual stdmet.
    today = dt.datetime.now(dt.timezone.utc).date()
    urls = []
    samples = []
    if end >= today - dt.timedelta(days=44):
        urls.append(f"{BASE}/data/realtime2/{station}.txt")
    for year in range(start.year, end.year+1):
        urls.append(f"{BASE}/view_text_file.php?filename={station.lower()}h{year}.txt.gz&dir=data/historical/stdmet/")
    errors = []
    for url in urls:
        try:
            r = session.get(url, timeout=45, headers={"User-Agent":"LIX-PSH-V2/0.3"})
            if r.status_code == 404: continue
            r.raise_for_status()
            content = gzip.decompress(r.content).decode("utf-8") if r.content[:2] == b"\x1f\x8b" else r.text
            for sample in parse_text(content, start, end):
                sample["url"] = r.url
                samples.append(sample)
        except (requests.RequestException, ValueError, OSError, UnicodeError) as exc:
            errors.append(f"{url}: {type(exc).__name__}")
    dedup={};conflicts=set()
    for sample in samples:
        t=sample['time']
        if t in dedup and any(dedup[t].get(k)!=sample.get(k) for k in ('wind','gust','pressure','dir')):
            conflicts.add(t)
        else:dedup[t]=sample
    for t in conflicts:
        dedup.pop(t,None);errors.append('Conflicting feeds at '+t.isoformat()+'; observation withheld')
    return sorted(dedup.values(), key=lambda x:x["time"]), urls, errors

def setpeak(sheet, r, col, peak, dircol=None):
    if peak is None: return
    value, sample = peak
    when = sample["time"]
    sheet.cell(r,col).value = round(value,1)
    if dircol and sample["dir"] is not None:
        sheet.cell(r,dircol).value = int(sample["dir"])
    offset = col+2 if dircol else col+1
    for j,part in enumerate((when.strftime("%H%M"),when.day,when.month,when.year)):
        sheet.cell(r,offset+j).value = part

def populate(workbook, qc, start, end, counts, audit=None):
    sheet = workbook["Wind and Pressure"]
    for row in range(2,sheet.max_row+1):
        site = identifier(sheet.cell(row,1).value)
        network = str(sheet.cell(row,8).value or "").strip().upper()
        station = station_id(sheet.cell(row,1))
        if not station or not site:
            continue
        try:
            obs,urls,errors=retrieve(station,start,end)
            def peak(field,reverse=True):
                good=[(v[field],v) for v in obs if v[field] is not None]
                return (max if reverse else min)(good,key=lambda x:x[0]) if good else None
            w,g,p=peak("wind"),peak("gust"),peak("pressure",False)
            # Keep IEM observations when both feeds describe the same linked AWOS.
            for field,col,dircol,extreme in (("wind",11,12,w),("gust",17,None,g),("pressure",23,None,p)):
                if extreme is None or sheet.cell(row,col).value is not None:continue
                if field == "gust":
                    setpeak(sheet,row,col,extreme,18)
                    # WDIR describes the mean wind, not a separately measured gust bearing.
                    sheet.cell(row,18).value=None
                else:setpeak(sheet,row,col,extreme,dircol)
                if audit:
                    value,sample=extreme
                    audit.add("Wind and Pressure",row,site,field,round(value,1),"hPa" if field=="pressure" else "kn",
                              sample["time"],sample["url"],raw_value=value if field=="pressure" else value/KNOTS_PER_MS,
                              raw_unit="hPa" if field=="pressure" else "m/s",details="NDBC header units verified; sampling and coverage require review")
            sheet.cell(row,28).value="I";sheet.cell(row,29).value="A"
            sheet.cell(row,30).value="NDBC sampling/exposure review; gust direction unavailable"
            collected=sum(x is not None for x in (w,g,p))
            status="COLLECTED" if collected==3 else "PARTIAL" if collected else "NO DATA"
            if errors and not obs: status="ERROR"
            counts["ndbc_"+status.lower().replace(" ","_")]+=1
            qc.append(["Wind and Pressure",site,network,status,
                       f"NDBC {station}; {len(obs)} observations; wind={w is not None}, gust={g is not None}, pressure={p is not None}; "+
                       ("; ".join(errors[:2]) if errors else ""),urls[0]])
        except Exception as exc:
            counts["ndbc_error"]+=1
            qc.append(["Wind and Pressure",site,network,"ERROR",str(exc)[:250],f"{BASE}/station_page.php?station={station.lower()}"])
        time.sleep(.1)
