"""NOAA CO-OPS water-level adapter; exact station hyperlinks and MHHW only."""
import datetime as dt
import re
import time
from urllib.parse import parse_qs, urlparse

import requests
from openpyxl.cell.cell import MergedCell
from common import finite, inventory, observation_now

BASE = "https://api.tidesandcurrents.noaa.gov/api/prod/datagetter"

def station_from_cell(cell):
    target = cell.hyperlink.target if cell.hyperlink else ""
    url = urlparse(target or "")
    if url.hostname not in ("tidesandcurrents.noaa.gov", "www.tidesandcurrents.noaa.gov"):
        return None
    sid = parse_qs(url.query).get("id", [""])[0]
    return sid if re.fullmatch(r"[0-9]{7}", sid) else None

DATUM_URL = "https://api.tidesandcurrents.noaa.gov/mdapi/prod/webapi/stations/{station}/datums.json"

def verify_datums(payload, station):
    reported = str(payload.get("id") or station)
    if reported != station: raise ValueError("Datum station identity mismatch")
    if not any(x.get("name") == "MHHW" and finite(x.get("value")) is not None
               for x in payload.get("datums", [])):
        raise ValueError("MHHW datum is not explicitly available for this station")
    return True

def datum_evidence(station, session=requests):
    r = session.get(DATUM_URL.format(station=station), params={"units":"english"}, timeout=20)
    r.raise_for_status()
    verify_datums(r.json(), station)
    return r.url

def collect(station, start, end, session=requests):
    all_values = []
    links = []
    cursor = start
    while cursor <= end:
        stop = min(cursor + dt.timedelta(days=29), end)
        params = {
            "product": "water_level", "application": "LIX_PSH_V2",
            "begin_date": cursor.strftime("%Y%m%d"),
            "end_date": stop.strftime("%Y%m%d"),
            "station": station, "datum": "MHHW", "units": "english",
            "time_zone": "gmt", "format": "json"
        }
        response = session.get(BASE, params=params, timeout=45)
        response.raise_for_status()
        links.append(response.url)
        payload = response.json()
        if payload.get("error"):
            raise ValueError("CO-OPS: " + str(payload["error"])[:170])
        reported = str((payload.get("metadata") or {}).get("id") or "")
        if reported and reported != station:
            raise ValueError("Station identity mismatch")
        for entry in payload.get("data", []):
            try:
                reading = finite(entry["v"], -30, 40)
                timestamp = dt.datetime.strptime(entry["t"], "%Y-%m-%d %H:%M").replace(tzinfo=dt.timezone.utc)
            except (KeyError, ValueError, TypeError):
                continue
            if reading is not None and start <= timestamp.date() <= end and timestamp<observation_now():
                entry = dict(entry, source_url=response.url)
                all_values.append((reading, timestamp, entry))
        cursor = stop + dt.timedelta(days=1)
    return (max(all_values, key=lambda x: x[0]) if all_values else None), links

def populate(workbook, qc, start, end, counts, audit=None):
    sheet = workbook["Water Level"]
    for row in range(2, sheet.max_row + 1):
        site_id = str(sheet.cell(row, 1).value or "").strip()
        station = station_from_cell(sheet.cell(row, 1))
        if not station or not site_id:
            continue
        source = str(sheet.cell(row, 13).value or "").strip().upper()
        if source != "NOS":
            qc.append(["Water Level", site_id, source, "CHECK SOURCE", "CO-OPS link but source M is not NOS", ""])
            continue
        for col in (7, 8, 9, 10, 11, 12, 14, 15):
            cell = sheet.cell(row, col)
            if not isinstance(cell, MergedCell):
                cell.value = None
        try:
            evidence = datum_evidence(station)
            peak, urls = collect(station, start, end)
            if peak is None:
                counts["coops_no_data"] += 1
                qc.append(["Water Level", site_id, "NOS", "NO DATA", "No valid MHHW readings", urls[0] if urls else ""])
                continue
            value, when, raw = peak
            if audit:
                audit.add("Water Level",row,site_id,"water",round(value,2),"ft",when,
                          raw["source_url"],datum="MHHW",evidence=evidence,raw_value=value,
                          details="CO-OPS q="+str(raw.get("q"))+" f="+str(raw.get("f")))
            sheet.cell(row, 7).value = round(value, 2)
            sheet.cell(row, 8).value = "MHHW"
            sheet.cell(row, 9).value = when.strftime("%H%M")
            sheet.cell(row, 10).value = when.day
            sheet.cell(row, 11).value = when.month
            sheet.cell(row, 12).value = when.year
            # Data flag semantics vary with product; flag for review rather than assuming verified.
            quality = str(raw.get("f") or "")
            sheet.cell(row, 14).value = "I"
            sheet.cell(row, 15).value = "CO-OPS data requires manual verification; quality flag: " + (quality or "unavailable")
            counts["coops_collected"] += 1
            qc.append(["Water Level", site_id, "NOS", "REVIEW REQUIRED",
                       f"{value:.2f} ft MHHW at {when:%Y-%m-%d %H:%M} UTC; f={quality or 'missing'}",
                       urls[0] if urls else ""])
        except Exception as exc:
            counts["coops_errors"] += 1
            qc.append(["Water Level", site_id, "NOS", "ERROR", type(exc).__name__+": datum/data request unsuccessful",
                       f"https://tidesandcurrents.noaa.gov/stationhome.html?id={station}"])
        time.sleep(0.1)
