"""Isolated XPTN-only historical source probe. No PSH production imports."""
import csv, datetime as dt, json, os, re, statistics
from pathlib import Path
from urllib.parse import urlparse
import requests

OUT=Path("xptn_probe_results")
OUT.mkdir(exist_ok=True)
START=dt.datetime(2026,10,8,tzinfo=dt.timezone.utc)
END=dt.datetime(2026,10,11,tzinfo=dt.timezone.utc)
SITE="XPTN"
results={"scope":"XPTN ONLY","period_utc":[START.isoformat(),END.isoformat()],"sources":{}}
def request(url,params=None,token=None,timeout=35):
    headers={"User-Agent":"PSH-Version2-XPTN-one-station-historical-prototype/0.1"}
    if token:headers["Authorization"]="Bearer "+token
    r=requests.get(url,params=params,headers=headers,timeout=timeout)
    return r
def dump(name,data):
    (OUT/name).write_text(json.dumps(data,indent=2,default=str,allow_nan=False))
def safe_http(source,r):
    return {"http_status":r.status_code,"content_type":r.headers.get("Content-Type",""),"bytes":len(r.content)}
def synopsis(obs):
    if not obs:return {"rows":0}
    dates=sorted(t for t in [x.get("time_utc") for x in obs] if t)
    delta=[(dt.datetime.fromisoformat(b)-dt.datetime.fromisoformat(a)).total_seconds()/60 for a,b in zip(dates,dates[1:])]
    return {"rows":len(obs),"first":dates[0] if dates else None,"last":dates[-1] if dates else None,"median_spacing_minutes":statistics.median(delta) if delta else None,"largest_gap_minutes":max(delta) if delta else None,"sample_observations":obs[:5],"max_wind_raw":max((o["wind"] for o in obs if isinstance(o.get("wind"),(int,float))),default=None),"max_gust_raw":max((o["gust"] for o in obs if isinstance(o.get("gust"),(int,float))),default=None)}
token=os.getenv("SYNOPTIC_TOKEN")
base="https://api.synopticdata.com/v2/stations"
for kind,uri,params in [
    ("metadata",base+"/metadata",{"stid":SITE,"token":token}),
    ("timeseries",base+"/timeseries",{"stid":SITE,"token":token,"start":"202610080000","end":"202610110000","vars":"wind_speed,wind_gust,wind_direction,sea_level_pressure,pressure","obtimezone":"utc","sensorvars":"1"})]:
    record={"requested":bool(token)}
    if token:
        try:
            r=request(uri,params)
            record.update(safe_http(kind,r))
            data=r.json()
            sm=data.get("SUMMARY",{})
            record["response_code"]=sm.get("RESPONSE_CODE")
            record["response_message"]=str(sm.get("RESPONSE_MESSAGE",""))[:120]
            stations=[s for s in data.get("STATION",[]) if s.get("STID","").upper()==SITE]
            record["matching_stations"]=len(stations)
            for s in stations:
                record["metadata"]={k:s.get(k) for k in ("STID","NAME","LATITUDE","LONGITUDE","ELEVATION","SENSOR_VARIABLES","PERIOD_OF_RECORD","NETWORK","STATUS") if k in s}
            record["units"]=data.get("UNITS")
            if kind=="timeseries":
                parsed=[]
                for s in stations:
                    ob=s.get("OBSERVATIONS",{})
                    for i,t in enumerate(ob.get("date_time",[])):
                        if not t:continue
                        try:t_utc=dt.datetime.fromisoformat(t.replace("Z","+00:00"))
                        except ValueError:continue
                        if not START<=t_utc<END:continue
                        item={"time_utc":t_utc.isoformat(),"wind":None,"gust":None,"direction":None,"pressure":None,"sensor_values":{}}
                        for k,v in ob.items():
                            if k=="date_time" or not isinstance(v,list):continue
                            if i>=len(v):continue
                            if k.startswith(("wind_speed_set_","wind_gust_set_","wind_direction_set_","sea_level_pressure_set_","pressure_set_")):
                                item["sensor_values"][k]=v[i]
                                typ="wind" if k.startswith("wind_speed_") else "gust" if k.startswith("wind_gust_") else "direction" if k.startswith("wind_direction_") else "pressure"
                                if item[typ] is None and isinstance(v[i],(float,int)):item[typ]=v[i]
                        parsed.append(item)
                record["observations"]=synopsis(parsed)
                if parsed:dump("synoptic_xptn_observations.json",parsed)
            if data.get("STATION"):record["returned_stids"]=[s.get("STID") for s in data.get("STATION",[])][:8]
        except Exception as exc:record["error_class"]=type(exc).__name__
    else:record["blocker"]="SYNOPTIC_TOKEN not configured for isolated workflow"
    results["sources"]["synoptic_"+kind]=record
# Legacy iWindsurf archive request — record response metadata only; do not claim parsed wind.
legacy_url="https://ww.weatherflow.com/windAndWhere.iws"
params={"Isection":"Wind Yesterday","days":"41","regionID":"115","siteID":"7532"}
try:
    r=request(legacy_url,params)
    record=safe_http("iwindsurf",r)
    record["redirect_hostname"]=urlparse(r.url).hostname
    record["page_title"]=(re.search(r"<title[^>]*>(.*?)</title>",r.text,re.I|re.S).group(1).strip()[:150] if re.search(r"<title[^>]*>(.*?)</title>",r.text,re.I|re.S) else None)
    record["mentions_XPTN"]="XPTN" in r.text.upper()
    record["mentions_7532"]="7532" in r.text
    record["visible_calendar_dates"]=sorted(set(re.findall(r"(?:10[/.-](?:0?[89]|10)[/.-]2026|2026[/.-]10[/.-](?:0?[89]|10))",r.text)))[:10]
    record["warning"]="Web page is not an authenticated observation API; no historical readings accepted unless independently parsed with timestamps/units"
    results["sources"]["iwindsurf_page"]=record
except Exception as exc:results["sources"]["iwindsurf_page"]={"error_class":type(exc).__name__,"requested":True}
# Tempest: only explicitly configured token; 150074 is UNVERIFIED and not treated as station identity.
wf_token=os.getenv("WEATHERFLOW_TOKEN") or os.getenv("TEMPEST_TOKEN")
for name,url in [("stations","https://swd.weatherflow.com/swd/rest/stations"),("claimed_station_150074","https://swd.weatherflow.com/swd/rest/stations/150074")]:
    if not wf_token:
        results["sources"]["tempest_"+name]={"attempted":False,"blocker":"No WEATHERFLOW_TOKEN or TEMPEST_TOKEN configured; cannot test owner-access to professional XPTN"}
        continue
    try:
        r=request(url,token=wf_token)
        record=safe_http(name,r)
        d=r.json()
        record["api_status"]=d.get("status")
        matches=[]
        for s in (d.get("stations") or d.get("locations") or []):
            if SITE in (str(s.get("name",""))+" "+str(s.get("public_name",""))).upper() or str(s.get("station_id"))=="150074":
                matches.append({k:s.get(k) for k in ("station_id","name","public_name","latitude","longitude")})
        record["matches"]=matches
        results["sources"]["tempest_"+name]=record
    except Exception as exc:results["sources"]["tempest_"+name]={"error_class":type(exc).__name__}
# Inspect only XPTN inventory and existing entries in the template; no cell changes.
try:
    from openpyxl import load_workbook
    wb=load_workbook("Copy of PSHLIX_YYYYALXX_StormName_Data.xlsx",read_only=True,data_only=True)
    sh=wb["Wind and Pressure"]
    rows=[]
    for row in sh.iter_rows(values_only=True):
        if "XPTN" in " ".join(str(v) for v in row[:8] if v is not None).upper():
            rows.append({"first_8_columns":[str(v) if v is not None else None for v in row[:8]],"wind":row[10],"gust":row[16],"pressure":row[22]})
    results["xptn_template_inventory"]=rows
except Exception as exc:results["template_error"]=type(exc).__name__
dump("xptn_source_report.json",results)
print("XPTN test complete:",{k:{kk:v.get(kk) for kk in ("http_status","response_code","matching_stations","blocker","error_class")} for k,v in results["sources"].items()})
